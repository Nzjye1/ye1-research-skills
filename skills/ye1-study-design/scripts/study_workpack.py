"""Portable design/CRF contract. Host supplies clinical judgment and confirmations.
Stdlib for work packages; openpyxl is lazy-loaded for XLSX. No cleaning/statistics.
"""
from __future__ import annotations

# All data-bearing streams stay behind the controlled local adapter.
if __name__ == '__main__':
    import sys as _privacy_sys
    if '--help' not in _privacy_sys.argv[1:]:
        from privacy_boundary import cli_local_entry
        cli_local_entry()

import argparse
import copy
import hashlib
import json
import math
import os
import re
import sys
from datetime import date, datetime, timezone
from pathlib import Path

SCHEMA = "ye1-study/1"
NAME = re.compile(r"^[a-z][a-z0-9_]{0,49}$")
CORE = ("case_kind", "plan", "dictionary", "collection", "ethics")
COVERS = ["population", "exposure", "outcomes", "time_anchors", "confounders", "analysis", "burden", "dictionary"]
DICT_HEADERS = ["表", "粒度", "字段", "中文名", "类型", "单位", "定义", "允许值/范围", "缺失码", "时间锚", "必要性", "来源", "角色"]
META_SHEETS = {"填写说明", "变量字典", "版本信息"}


def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def canonical(obj):
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def fingerprint(obj):
    return hashlib.sha256(canonical(obj).encode()).hexdigest()


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path, obj):
    Path(path).write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def require(condition, message):
    if not condition:
        raise ValueError(message)


def core_hash(s):
    return fingerprint({k: s[k] for k in CORE})


def save_ledger(root, ledger):
    path = Path(root) / "work-package.json"
    temp = path.with_suffix(".tmp")
    write_json(temp, ledger)
    temp.replace(path)


def validate_spec(s):
    require(s.get("schema_version") == SCHEMA, "unsupported schema_version")
    require(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}", s.get("study_id", "")), "invalid study_id")
    require(s.get("case_kind") in {"SYNTHETIC", "REAL_STUDY"} and s.get("title"), "study identity required")
    for k in ("question", "design", *COVERS[:-1], "sample_size"):
        require(s.get("plan", {}).get(k), "plan requires " + k)
    sample = s["plan"]["sample_size"]
    require(sample.get("status") in {"OPEN", "ASSUMPTION", "CONFIRMED"} and sample.get("rationale"), "sample size assumptions/status required")
    c = s.get("collection", {})
    require(c.get("mode") in {"RETROSPECTIVE", "FUTURE_ROUTINE", "MIXED"}, "collection mode required")
    for k in ("setting", "source", "recording_method", "staffing", "period", "known_unavailable"):
        require(k in c, "collection requires " + k)
    follow = c.get("followup", {})
    require(type(follow.get("included")) is bool, "explicit followup decision required")
    if follow["included"]:
        require(follow.get("capability_confirmed") is True and follow.get("doctor_agreed") is True
                and follow.get("source") and follow.get("schedule"), "followup needs capacity and doctor agreement")
    for k in ("review", "institution", "consent", "local_data_use"):
        e = s.get("ethics", {}).get(k, {})
        require(e.get("status") in {"NOT_PROVIDED", "PENDING", "VERIFIED", "EXEMPT_VERIFIED", "SIMULATED_ONLY"} and e.get("note"), "ethics status/note required: " + k)
        if e["status"] in {"VERIFIED", "EXEMPT_VERIFIED"}:
            require(e.get("source") and e.get("date"), "verified ethics requires evidence source/date")
        require(not (s["case_kind"] == "REAL_STUDY" and e["status"] == "SIMULATED_ONLY"), "simulation cannot authorize real research")
    for f in s.get("facts", []):
        require(f.get("status") in {"CONFIRMED", "ASSUMPTION", "OPEN"} and f.get("source") and f.get("date") and f.get("key"), "facts need status/source/date/key")
    for f in s.get("findings", []):
        require(f.get("severity") in {"CRITICAL", "ADVISORY"} and f.get("status") in {"OPEN", "UNAVAILABLE", "RESOLVED"}
                and f.get("rationale") and f.get("evidence_refs"), "finding needs reason/evidence/severity/status")
        require(f.get("affects") and set(f["affects"]) <= {"QC", "COLLECTION", "DESCRIPTIVE", "CONFIRMATORY", "PUBLICATION"}, "invalid finding activity")
    tables = s.get("dictionary", {}).get("tables", [])
    require(tables, "data tables required")
    names = [t["name"] for t in tables]
    require(len(set(names)) == len(names), "duplicate tables")
    for t in tables:
        require(NAME.fullmatch(t["name"]) and len(t["name"]) <= 31 and t.get("grain") and t.get("label"), "table name/grain/label required")
        fields = t.get("fields", [])
        fieldnames = [f["name"] for f in fields]
        require(fields and len(set(fieldnames)) == len(fields), "empty/duplicate fields")
        require(t.get("primary_key") and set(t["primary_key"]) <= set(fieldnames), "invalid primary key")
        for f in fields:
            require(NAME.fullmatch(f["name"]) and "__" not in f["name"], "invalid/reserved field name")
            require(f["name"] not in {"patient_name", "phone", "mobile", "national_id", "medical_record_number"}
                    and f.get("role") != "direct_identifier", "shared CRF requires deidentified IDs")
            require(f.get("type") in {"string", "integer", "number", "category", "date", "datetime"} and type(f.get("required")) is bool, "invalid type/required")
            for k in ("label", "definition", "unit", "time_anchor", "source", "role"):
                require(isinstance(f.get(k), str) and f[k], "field requires " + k)
            codes = f.get("missing_codes")
            require(isinstance(codes, dict) and all(NAME.fullmatch(k.lower()) and isinstance(v, str) and v for k, v in codes.items()), "missing codes must map text to definitions")
            if f["name"] in t["primary_key"]:
                require(f["required"] and not codes, "keys required without missing codes")
            if f["type"] == "category":
                vals = f.get("allowed_values")
                require(vals and all(isinstance(v, str) for v in vals) and len(set(vals)) == len(vals), "explicit unique text categories required")
            if "pattern" in f:
                re.compile(f["pattern"])
            if "min" in f and "max" in f:
                require(f["min"] <= f["max"], "invalid bounds")
        for fk in t.get("foreign_keys", []):
            require(fk.get("table") in names and set(fk.get("fields", [])) <= set(fieldnames), "invalid foreign key")
            parent = next(x for x in tables if x["name"] == fk["table"])
            require(fk.get("references") == parent["primary_key"] and len(fk["fields"]) == len(fk["references"]), "foreign key must reference full primary key")
        for r in t.get("constraints", []):
            require(r.get("kind") in {"order", "window"} and r.get("left") in fieldnames and r.get("right") in fieldnames and r.get("message"), "invalid constraint")
            if r["kind"] == "order":
                require(r.get("op") in {"<=", ">=", "<", ">", "=="}, "invalid comparison")
            else:
                require(r.get("min") is not None and r.get("max") is not None and r["min"] <= r["max"], "window bounds required")
        if "longitudinal" in t:
            r = t["longitudinal"]
            require(r.get("group_by") and set(r["group_by"]) <= set(fieldnames) and r.get("start") in fieldnames
                    and r.get("end") in fieldnames, "invalid longitudinal keys/time")
            if r.get("terminal_values"):
                require(r.get("terminal_field") in fieldnames, "terminal event field missing")
    canonical(s)
    return s


def load_package(root):
    root = Path(root)
    ledger = read_json(root / "work-package.json")
    require(ledger.get("schema_version") == SCHEMA, "unsupported work package")
    for item in ledger["versions"]:
        expected = f"versions/v{item['revision']:04d}/study.json"
        require(item["path"] == expected and file_hash(root / expected) == item["sha256"], "snapshot path/hash changed")
        saved = validate_spec(read_json(root / expected))
        require(saved["study_id"] == ledger["study_id"], "study identity conflict")
    active = next(x for x in ledger["versions"] if x["revision"] == ledger["active_revision"])
    return ledger, read_json(root / active["path"])


def confirmation(ledger, spec):
    kind = "SIMULATED" if spec["case_kind"] == "SYNTHETIC" else "DOCTOR"
    return next((c for c in reversed(ledger["confirmations"]) if c["core_sha256"] == core_hash(spec) and c["kind"] == kind), None)


def blocked_activities(spec):
    return {a for f in spec.get("findings", []) if f["severity"] == "CRITICAL" and f["status"] != "RESOLVED" for a in f["affects"]}


def collection_readiness(ledger, spec):
    if spec["case_kind"] == "SYNTHETIC":
        return "SIMULATION_ONLY" if confirmation(ledger, spec) else "SIMULATION_DRAFT"
    verified = all(spec["ethics"][k]["status"] in {"VERIFIED", "EXEMPT_VERIFIED"} for k in ("review", "institution", "consent", "local_data_use"))
    return "HOST_DOCUMENTED_PREREQUISITES" if verified and confirmation(ledger, spec) and "COLLECTION" not in blocked_activities(spec) else "DRAFT_NOT_AUTHORIZED_FOR_REAL_COLLECTION"


def readiness_label(ledger, spec):
    return {"SIMULATION_ONLY": "仅用于合成演示", "SIMULATION_DRAFT": "合成方案草案，待模拟确认",
            "HOST_DOCUMENTED_PREREQUISITES": "宿主已记录前提证据，仍以机构实际批准范围为准",
            "DRAFT_NOT_AUTHORIZED_FOR_REAL_COLLECTION": "研究草案；真实收集前提尚未核清"}[collection_readiness(ledger, spec)]


def columns(table):
    result = []
    for f in table["fields"]:
        label = f["label"] + (" (" + f["unit"] + ")" if f["unit"] != "无" else "") + (" *" if f["required"] else "") + "\n" + f["name"]
        result.append((f["name"], label, f, False))
        if f["missing_codes"]:
            result.append((f["name"] + "__missing", "缺失原因\n" + f["name"] + "__missing", f, True))
    return result


def dictionary_rows(spec):
    rows = []
    for t in spec["dictionary"]["tables"]:
        for f in t["fields"]:
            allowed = {k: f[k] for k in ("allowed_values", "value_labels", "min", "max", "pattern") if k in f}
            rows.append([t["name"], t["grain"], f["name"], f["label"], f["type"], f["unit"], f["definition"],
                         canonical(allowed), canonical(f["missing_codes"]), f["time_anchor"], "必要" if f["required"] else "可选", f["source"], f["role"]])
    return rows


def prose(value):
    return value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, indent=2)


def protocol_text(spec):
    lines = ["# " + spec["title"], "", "研究ID：" + spec["study_id"], "身份：" + spec["case_kind"], ""]
    labels = ["研究问题", "设计", "人群", "暴露", "结局", "时间锚", "混杂与偏倚", "分析思路", "收集负担", "样本量及假设"]
    keys = ["question", "design", *COVERS[:-1], "sample_size"]
    for key, label in zip(keys, labels):
        lines += ["## " + label, "", prose(spec["plan"][key]), ""]
    for key, label in (("facts", "已答情况"), ("open_items", "待办"), ("ethics", "伦理与机构情况")):
        lines += ["## " + label, "", prose(spec.get(key, [])), ""]
    return "\n".join(lines)


def collection_text(ledger, spec):
    lines = ["# 填写与续接说明", "", f"研究 {spec['study_id']}；方案 v{ledger['active_revision']:04d}",
             "当前收集前提：" + readiness_label(ledger, spec),
             "CRF不授予真实招募或数据使用批准；真实研究需机构及伦理适用要求的实际证据。", "",
             "只用去标识ID。姓名、手机号、身份证、住院号和编码对应表留在机构获准环境，不放共享CRF。",
             "表头有ID、中文名、单位；* 表示必要。数值格只填数值，不带单位。",
             "缺失时值格留空、相邻__missing栏选原因；不能同时填值和缺失码，0不代表缺失。",
             "日期YYYY-MM-DD；小时或日偏移遵守字典时间锚，不能把日数填小时列。",
             "粘贴能绕过Excel校验，回传必须再QC；不自动删除、纠正、換算或插补。", ""]
    for t in spec["dictionary"]["tables"]:
        lines += [f"- {t['name']}（{t['label']}）：{t['grain']}；唯一键：{', '.join(t['primary_key'])}。"]
        lines += ["  约束：" + r["message"] for r in t.get("constraints", [])]
    lines += ["", "收集方式：" + prose(spec["collection"]), "",
              "回传保留整份Excel与工作包目录。填表另存副本，保留版本信息、变量字典页。",
              "等待收数是正常进度，没有后台催收或预约；下次读回本目录即可继续。",
              "改单项记录理由、来源与日期。终点、时点、人群或方法变更重新确认，旧版保留。",
              "样本量/精度目标未确认就留待办；预留行数只是填写容量。"]
    return "\n".join(lines) + "\n"


def save_version(root, ledger, spec, revision):
    out = Path(root) / "versions" / f"v{revision:04d}"
    out.mkdir(parents=True, exist_ok=False)
    write_json(out / "study.json", spec)
    write_json(out / "dictionary.json", spec["dictionary"])
    (out / "研究方案.md").write_text(protocol_text(spec), encoding="utf-8")
    ledger["active_revision"] = revision
    ledger["versions"].append({"revision": revision, "path": f"versions/v{revision:04d}/study.json", "sha256": file_hash(out / "study.json")})


def create_package(spec, out):
    validate_spec(spec)
    Path(out).mkdir(parents=True, exist_ok=False)
    ledger = {"schema_version": SCHEMA, "study_id": spec["study_id"], "active_revision": 1, "versions": [], "confirmations": [],
              "events": [], "progress": {"status": "DESIGN_DISCUSSION", "note": "待讨论确认", "updated_at": now()}}
    save_version(out, ledger, spec, 1)
    save_ledger(out, ledger)
    return resume(out)


def confirm_package(root, record):
    ledger, spec = load_package(root)
    require(record.get("study_id") == spec["study_id"] and record.get("revision") == ledger["active_revision"], "confirmation belongs to another study/version")
    require(record.get("core_sha256") == core_hash(spec), "confirmation must bind current core_sha256")
    require(record.get("kind") == ("SIMULATED" if spec["case_kind"] == "SYNTHETIC" else "DOCTOR"), "simulation/real confirmation mismatch")
    for k in ("by", "date", "source", "quote"):
        require(record.get(k), "confirmation needs " + k)
    require(record.get("covers") == COVERS, "confirmation must cover design, definitions and burden")
    from confirmation_evidence import validate as validate_confirmation
    source_evidence=validate_confirmation(record,simulated=record["kind"]=="SIMULATED")
    record={**record,"source_evidence":source_evidence}
    ledger["confirmations"].append(copy.deepcopy(record))
    ledger["events"].append({"type": "CONFIRMATION_RECORDED", "date": now(), "revision": ledger["active_revision"], "source": record["source"]})
    save_ledger(root, ledger)
    return resume(root)


def amend(root, change):
    ledger, old = load_package(root)
    require(change.get("base_revision") == ledger["active_revision"], "stale amendment base")
    require(all(change.get(k) for k in ("reason", "source", "date", "replacements")), "amendment needs reason/source/date/replacements")
    spec, diff = copy.deepcopy(old), []
    for path, value in change["replacements"].items():
        parts = path.strip("/").split("/")
        require(path.startswith("/") and parts[0] in {*CORE, "facts", "open_items", "title", "evidence", "findings", "proposal_options"} and parts[0] != "case_kind", "unsupported amendment path")
        node = spec
        for key in parts[:-1]:
            node = node[int(key)] if isinstance(node, list) else node[key]
        key = int(parts[-1]) if isinstance(node, list) else parts[-1]
        previous = copy.deepcopy(node[key])
        node[key] = copy.deepcopy(value)
        diff.append({"path": path, "before": previous, "after": value})
    validate_spec(spec)
    require(spec != old, "no actual amendment")
    revision = ledger["active_revision"] + 1
    save_version(root, ledger, spec, revision)
    ledger["events"].append({"type": "AMENDMENT", "revision": revision, "date": change["date"], "source": change["source"],
                             "reason": change["reason"], "changes": diff, "scientific_confirmation_invalidated": core_hash(old) != core_hash(spec)})
    if core_hash(old) != core_hash(spec):
        ledger["progress"] = {"status": "DESIGN_DISCUSSION", "note": "确认变更项及影响，既有答复保留", "updated_at": now()}
    save_ledger(root, ledger)
    return resume(root)


def record_progress(root, record):
    ledger, _ = load_package(root)
    require(record.get("status") in {"DESIGN_DISCUSSION", "WAITING_FOR_COLLECTION", "COLLECTING", "RETURNED"}, "unknown progress")
    require(all(record.get(k) for k in ("note", "date", "source")), "progress needs note/date/source")
    ledger["progress"] = copy.deepcopy(record)
    ledger["events"].append({"type": "PROGRESS", **record})
    save_ledger(root, ledger)
    return resume(root)


def resume(root):
    ledger, s = load_package(root)
    c = confirmation(ledger, s)
    return {"schema_version": SCHEMA, "study_id": s["study_id"], "title": s["title"], "case_kind": s["case_kind"],
            "revision": ledger["active_revision"], "core_sha256": core_hash(s), "scientific_confirmation_valid": c is not None,
            "confirmation_kind": c["kind"] if c else None, "collection_readiness": collection_readiness(ledger, s),
            "progress": ledger["progress"], "known_unavailable": s["collection"]["known_unavailable"],
            "facts": s.get("facts", []), "open_items": s.get("open_items", []), "plan": s["plan"],
            "dictionary": s["dictionary"], "ethics": s["ethics"], "events": ledger["events"], "automatic_followup_scheduled": False}


def meta_rows(ledger, s):
    return [["schema_version", SCHEMA], ["study_id", s["study_id"]], ["revision", ledger["active_revision"]],
            ["spec_sha256", fingerprint(s)], ["dictionary_sha256", fingerprint(s["dictionary"])],
            ["core_sha256", core_hash(s)], ["case_kind", s["case_kind"]]]


def export_crf(root, out, rows=200, text_only=False):
    ledger, s = load_package(root)
    require(type(rows) is int and 1 <= rows <= 10000, "capacity 1..10000 required, not sample size")
    target = Path(out)
    target.mkdir(parents=True, exist_ok=False)
    write_json(target / "dictionary.json", s["dictionary"])
    (target / "研究方案.md").write_text(protocol_text(s), encoding="utf-8")
    (target / "填写说明.md").write_text(collection_text(ledger, s), encoding="utf-8")
    write_json(target / "续接记录.json", resume(root))
    (target / "变量字典.md").write_text("# 变量字典\n\n" + prose(s["dictionary"]) + "\n", encoding="utf-8")
    status = "TEXT_ONLY_REQUESTED" if text_only else "XLSX_CREATED"
    if not text_only:
        try:
            import openpyxl
        except ImportError:
            status = "TEXT_ONLY_OPENPYXL_UNAVAILABLE"
        else:
            build_workbook(ledger, s, target / "空白CRF.xlsx", rows)
    files = {p.name: file_hash(p) for p in target.iterdir() if p.is_file()}
    result = {"status": status, "study_id": s["study_id"], "revision": ledger["active_revision"],
              "scientific_confirmation_valid": bool(confirmation(ledger, s)), "collection_readiness": collection_readiness(ledger, s),
              "files": files, "row_capacity_is_not_sample_size": True, "created_at": now()}
    write_json(target / "export-receipt.json", result)
    return result


def build_workbook(ledger, s, path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    guide = wb.create_sheet("填写说明")
    for row in [
        ["Ye1 CRF", s["title"]],
        ["研究与方案", f"{s['study_id']} / v{ledger['active_revision']:04d}"],
        ["收集前提", readiness_label(ledger, s)],
        ["身份", "纯合成演示，不是真实招募许可" if s["case_kind"] == "SYNTHETIC" else "真实收集需适用机构与伦理实际批准"],
        ["确认", ("已记录模拟确认" if s["case_kind"] == "SYNTHETIC" else "已记录对应方案确认") if confirmation(ledger, s) else "尚未确认，供讨论"],
        ["填写规则", "* 为必要项；浅黄色格可填。数值不带单位；缺失时值留空并填相邻缺失原因。"],
        ["隐私", "仅去标识ID；姓名电话与对应表留在机构获准环境，不放此表。"],
        ["容量", f"每张表预留{rows}行，仅填写容量，不代表样本量；超出后需扩展校验范围。"],
        ["回传", "保留版本信息、变量字典与整份工作包；粘贴可绕过Excel校验，回传再QC。"],
        ["等待", "正常等待收数，无后台自动追问。"],
    ]:
        guide.append(row)
    for t in s["dictionary"]["tables"]:
        guide.append([t["label"], t["grain"] + "；唯一键 " + ", ".join(t["primary_key"])])
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 100
    for row in guide:
        guide.row_dimensions[row[0].row].height = 38
    for t in s["dictionary"]["tables"]:
        ws = wb.create_sheet(t["name"])
        cols = columns(t)
        ws.append([x[1] for x in cols])
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{rows+1}"
        ws.row_dimensions[1].height = 48
        for col, (field_name, _, f, missing) in enumerate(cols, 1):
            letter = get_column_letter(col)
            ws.column_dimensions[letter].width = max(25, min(55, len(field_name) * 1.12))
            cellref, end = f"{letter}2", f"{letter}{rows+1}"
            choices = list(f["missing_codes"]) if missing else f.get("allowed_values")
            if choices and sum(len(v) + 1 for v in choices) < 250 and all('"' not in v and ',' not in v for v in choices):
                dv = DataValidation(type="list", formula1='"' + ','.join(choices) + '"', allow_blank=True)
            else:
                tests = [f'ISNUMBER({cellref})'] if f["type"] in {"integer", "number", "date", "datetime"} else [f'ISTEXT({cellref})']
                if f["type"] == "integer":
                    tests.append(f'MOD({cellref},1)=0')
                if f["type"] in {"number", "integer"}:
                    for key, op in (("min", ">="), ("max", "<=")):
                        if key in f:
                            tests.append(f'{cellref}{op}{f[key]}')
                dv = DataValidation(type="custom", formula1=f'OR({cellref}="",AND(' + ','.join(tests) + '))', allow_blank=True)
            dv.showErrorMessage, dv.showInputMessage = True, True
            dv.errorTitle, dv.error = "请按字典填写", "检查类型、允许值、范围与单位；缺失原因另填。"
            dv.promptTitle = f["label"][:32]
            dv.prompt = ("缺失码：" + canonical(f["missing_codes"]) if missing else f["definition"] + "；" + f["time_anchor"])[:250]
            ws.add_data_validation(dv)
            dv.add(f"{cellref}:{end}")
            for row in range(2, rows + 2):
                cell = ws.cell(row, col)
                cell.fill = PatternFill("solid", fgColor="FFF8E4")
                cell.number_format = "@" if missing or f["type"] in {"string", "category"} else {"date": "yyyy-mm-dd", "datetime": "yyyy-mm-dd hh:mm", "integer": "0", "number": "0.00"}[f["type"]]
            if missing:
                prev = get_column_letter(col - 1)
                formula = 'AND($' + prev + '2<>"",$' + letter + '2<>"")'
                ws.conditional_formatting.add(f"{prev}2:{letter}{rows+1}", FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor="FCE4D6")))
        # Cross-field time constraints are also visible while typing. QC is authoritative.
        index = {name: get_column_letter(i) for i, (name, _, _, _) in enumerate(cols, 1)}
        for rule in t.get("constraints", []):
            left, right = "$" + index[rule["left"]] + "2", "$" + index[rule["right"]] + "2"
            if rule["kind"] == "order":
                relation = left + ("=" if rule["op"] == "==" else rule["op"]) + right
            else:
                factor = "*24" if rule.get("unit") == "hours" and next(f for f in t["fields"] if f["name"] == rule["left"])["type"] in {"date", "datetime"} else ""
                delta = f"({left}-{right}){factor}"
                relation = f"AND({delta}>={rule['min']},{delta}<={rule['max']})"
            formula = f'AND(ISNUMBER({left}),ISNUMBER({right}),NOT({relation}))'
            ws.conditional_formatting.add(f"A2:{get_column_letter(len(cols))}{rows+1}", FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor="FCE4D6")))
    dictionary = wb.create_sheet("变量字典")
    dictionary.append(DICT_HEADERS)
    for row in dictionary_rows(s):
        dictionary.append(row)
    dictionary.freeze_panes = "D2"
    for i in range(1, len(DICT_HEADERS) + 1):
        dictionary.column_dimensions[get_column_letter(i)].width = 24 if i not in (7, 8, 9, 10) else 48
    dictionary.row_dimensions[1].height = 32
    for row in range(2, dictionary.max_row + 1):
        dictionary.row_dimensions[row].height = 68
    meta = wb.create_sheet("版本信息")
    meta.append(["字段", "值"])
    for row in meta_rows(ledger, s):
        meta.append(row)
    meta.column_dimensions["A"].width, meta.column_dimensions["B"].width = 26, 75
    for ws in wb:
        ws.sheet_view.showGridLines = False
        for row in ws:
            for cell in row:
                cell.font = Font(name="Arial", size=11, bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "243146")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="243B5A")
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # host text is never executed as a formula
    wb.save(path)


def value_ok(v, f):
    typ = f["type"]
    if typ in {"integer", "number"}:
        if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v):
            return False
        if typ == "integer" and v != int(v):
            return False
        if ("min" in f and v < f["min"]) or ("max" in f and v > f["max"]):
            return False
    elif typ in {"date", "datetime"}:
        if not isinstance(v, (date, datetime)):
            return False
        if typ == "date" and isinstance(v, datetime) and v.time().isoformat() != "00:00:00":
            return False
    elif not isinstance(v, str) or not v.strip():
        return False
    if "allowed_values" in f and v not in f["allowed_values"]:
        return False
    if "pattern" in f and not re.fullmatch(f["pattern"], str(v)):
        return False
    return True


def qc_workbook(root, workbook):
    from openpyxl import load_workbook
    ledger, s = load_package(root)
    # Real-data access permission is checked before opening the workbook.
    if s["case_kind"] == "REAL_STUDY":
        require(s["ethics"]["local_data_use"]["status"] in {"VERIFIED", "EXEMPT_VERIFIED"} and "QC" not in blocked_activities(s),
                "real-data access not verified; keep draft and obtain institution-authorized scope before reading records")
    wb = load_workbook(workbook, data_only=False)
    issues, counts, missingness, records, keys = [], {}, {}, {}, {}
    def add(code, message, table=None, row=None, field=None, severity="CRITICAL"):
        issues.append({"code": code, "severity": severity, "message": message, "table": table, "row": row, "field": field})
    if "版本信息" not in wb.sheetnames:
        add("METADATA_MISSING", "缺少版本信息；先核来源，不自动套当前字典")
    else:
        actual = {r[0]: r[1] for r in wb["版本信息"].iter_rows(min_row=2, values_only=True)}
        for key, value in meta_rows(ledger, s):
            if actual.get(key) != value:
                add("VERSION_MISMATCH", "研究/定义/版本不一致：" + key)
    if "变量字典" not in wb.sheetnames or [list(r) for r in wb["变量字典"].values] != [DICT_HEADERS] + dictionary_rows(s):
        add("DICTIONARY_CHANGED", "字典缺失或定义/单位被更改；需确认变更并新建版本")
    if set(wb.sheetnames) - (META_SHEETS | {t["name"] for t in s["dictionary"]["tables"]}):
        add("UNEXPECTED_SHEET", "未定义工作表；不要默默合并或分享其中身份资料")
    for t in s["dictionary"]["tables"]:
        name = t["name"]
        records[name], keys[name], missingness[name], counts[name] = [], {}, {}, 0
        if name not in wb.sheetnames:
            add("TABLE_MISSING", "缺少数据表", name)
            continue
        ws, cols = wb[name], columns(t)
        if [c.value for c in ws[1]] != [x[1] for x in cols]:
            add("HEADER_CHANGED", "表头/单位/列数不符；不猜测映射", name)
            continue
        for cells in ws.iter_rows(min_row=2):
            vals = [c.value for c in cells]
            if all(v is None for v in vals):
                continue
            rownum, row = cells[0].row, dict(zip([x[0] for x in cols], vals))
            counts[name] += 1
            records[name].append((rownum, row))
            for cell, (field, _, _, _) in zip(cells, cols):
                if cell.data_type in {"f", "e"}:
                    add("FORMULA_OR_ERROR", "原始CRF不接受公式/错误值，请核实实际记录", name, rownum, field)
            for f in t["fields"]:
                field = f["name"]
                v, code = row[field], row.get(field + "__missing")
                m = missingness[name].setdefault(field, {"missing": 0, "rows": 0, "codes": {}})
                m["rows"] += 1
                if v is None:
                    m["missing"] += 1
                    if code in f["missing_codes"]:
                        m["codes"][code] = m["codes"].get(code, 0) + 1
                        add("DOCUMENTED_MISSING", "缺失已记原因，按方案定分析子集；不自动插补", name, rownum, field, "ADVISORY")
                    elif code is not None:
                        add("MISSING_CODE_INVALID", "未知缺失码", name, rownum, field)
                    elif f["required"]:
                        add("REQUIRED_MISSING", "必要项缺失且未记原因", name, rownum, field)
                    else:
                        add("OPTIONAL_BLANK", "可选项未填，确认是否未采集", name, rownum, field, "ADVISORY")
                else:
                    if code is not None:
                        add("VALUE_AND_MISSING_CODE", "已有值同时标缺失", name, rownum, field)
                    if not value_ok(v, f):
                        add("INVALID_VALUE", "类型/范围/编码不符合字典", name, rownum, field)
            key = tuple(row[f] for f in t["primary_key"])
            if all(v is not None for v in key):
                if key in keys[name]:
                    add("DUPLICATE_KEY", "同一记录键重复；与允许的一人多次记录区分", name, rownum)
                else:
                    keys[name][key] = rownum
            definitions = {f["name"]: f for f in t["fields"]}
            for r in t.get("constraints", []):
                left, right = row[r["left"]], row[r["right"]]
                if left is None or right is None or not value_ok(left, definitions[r["left"]]) or not value_ok(right, definitions[r["right"]]):
                    continue
                try:
                    if r["kind"] == "order":
                        ok = {"<=": lambda: left <= right, ">=": lambda: left >= right, "<": lambda: left < right,
                              ">": lambda: left > right, "==": lambda: left == right}[r["op"]]()
                    else:
                        delta = left - right
                        if hasattr(delta, "total_seconds"):
                            delta = delta.total_seconds() / {"hours": 3600, "days": 86400}[r.get("unit", "days")]
                        ok = r["min"] <= delta <= r["max"]
                    if not ok:
                        add("TIME_OR_ORDER", r["message"], name, rownum)
                except (TypeError, KeyError):
                    add("CONSTRAINT_TYPE", "约束字段类型或单位需核", name, rownum)
    for t in s["dictionary"]["tables"]:
        for fk in t.get("foreign_keys", []):
            for rownum, row in records[t["name"]]:
                if tuple(row[f] for f in fk["fields"]) not in keys[fk["table"]]:
                    add("ORPHAN_RECORD", "关联键未出现在主表，不能合并", t["name"], rownum)
        rule = t.get("longitudinal")
        if rule:
            groups = {}
            definitions = {f["name"]: f for f in t["fields"]}
            for rownum, row in records[t["name"]]:
                if not all(row[k] is not None for k in rule["group_by"]):
                    continue
                if not all(row[k] is not None and value_ok(row[k], definitions[k]) for k in (rule["start"], rule["end"])):
                    continue
                groups.setdefault(tuple(row[k] for k in rule["group_by"]), []).append((rownum, row))
            for grouped in groups.values():
                grouped.sort(key=lambda item: item[1][rule["start"]])
                latest_end, terminal = None, False
                for rownum, row in grouped:
                    start, end = row[rule["start"]], row[rule["end"]]
                    if latest_end is not None and start < latest_end:
                        add("EPISODE_OVERLAP", "同一对象记录时间重叠，核对时点及记录粒度", t["name"], rownum)
                    if terminal:
                        add("AFTER_TERMINAL_EVENT", "终止事件后仍有记录，核对身份与时间", t["name"], rownum)
                    latest_end = end if latest_end is None else max(latest_end, end)
                    terminal = terminal or row.get(rule.get("terminal_field")) in rule.get("terminal_values", [])
    if not any(counts.values()):
        add("EMPTY_RETURN", "空白为正常待收数，尚无可交分析记录", severity="ADVISORY")
    c = confirmation(ledger, s)
    if not c:
        add("CONFIRMATION_MISSING", "当前科学定义无有效确认；确认变更项及影响")
    critical = any(x["severity"] == "CRITICAL" for x in issues)
    blocked = blocked_activities(s)
    eligible = bool(any(counts.values()) and not critical and "QC" not in blocked)
    report = {"schema_version": SCHEMA, "study_id": s["study_id"], "revision": ledger["active_revision"], "case_kind": s["case_kind"],
              "source_sha256": file_hash(workbook), "core_sha256": core_hash(s), "dictionary_sha256": fingerprint(s["dictionary"]),
              "checked_at": now(), "row_counts": counts, "missingness": missingness, "issues": issues,
              "qc_pass": not critical, "analysis_handoff_eligible": eligible, "scientific_confirmation_valid": bool(c),
              "mode": "SYNTHETIC_DEMO" if s["case_kind"] == "SYNTHETIC" else "REAL_STUDY",
              "blocked_activities": sorted(blocked), "cleaning_performed": False, "statistics_performed": False}
    wb.close()
    return report


def return_and_handoff(root, workbook, out):
    ledger, s = load_package(root)
    report = qc_workbook(root, workbook)
    target = Path(out)
    target.mkdir(parents=True, exist_ok=False)
    write_json(target / "qc-report.json", report)
    lines = ["# 回传核对", "", f"研究 {s['study_id']}；方案 v{ledger['active_revision']:04d}",
             f"QC通过：{report['qc_pass']}；可交分析入口：{report['analysis_handoff_eligible']}", "未清洗、插补或计算统计结果。", ""]
    lines += [f"- {i['severity']} / {i['code']}：{i['message']}（{i['table']}，行{i['row']}，{i['field']}）" for i in report["issues"]]
    (target / "回传核对.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    h = {"schema_version": "ye1-analysis-handoff/1", "study_id": s["study_id"], "revision": ledger["active_revision"],
         "case_kind": s["case_kind"], "eligible": report["analysis_handoff_eligible"],
         "source": {"path": os.path.relpath(Path(workbook).resolve(), target.resolve()), "sha256": report["source_sha256"]},
         "work_package": os.path.relpath(Path(root).resolve(), target.resolve()),
         "study_sha256": fingerprint(s), "core_sha256": core_hash(s), "dictionary_sha256": report["dictionary_sha256"],
         "scientific_confirmation": confirmation(ledger, s), "plan": s["plan"], "dictionary": s["dictionary"],
         "collection": s["collection"], "ethics": s["ethics"], "findings": s.get("findings", []), "evidence": s.get("evidence", []),
         "open_items": s.get("open_items", []), "qc_report": "qc-report.json", "qc_report_sha256": file_hash(target / "qc-report.json"),
         "next_entry": "Ye1-数据质控 → Ye1-统计分析方案/对应已确认分析", "requires_downstream_scientific_method_check": True,
         "automatic_analysis_performed": False, "submission_ready": False}
    write_json(target / "analysis-handoff.json", h)
    return report


def validate_handoff(path, source_override=None, package_override=None):
    p, h = Path(path), read_json(path)
    require(h.get("schema_version") == "ye1-analysis-handoff/1" and h.get("eligible") is True, "handoff not eligible")
    source = Path(source_override) if source_override else p.parent / h["source"]["path"]
    package = Path(package_override) if package_override else p.parent / h["work_package"]
    ledger, active = load_package(package)
    require(ledger["active_revision"] == h["revision"] and fingerprint(active) == h["study_sha256"]
            and active["study_id"] == h["study_id"], "handoff is stale against current work package")
    require(confirmation(ledger, active) == h["scientific_confirmation"], "active confirmation changed")
    require(file_hash(source) == h["source"]["sha256"], "returned data changed after QC")
    require(h["qc_report"] == "qc-report.json" and file_hash(p.parent / h["qc_report"]) == h["qc_report_sha256"], "QC report changed")
    q = read_json(p.parent / h["qc_report"])
    require(q["analysis_handoff_eligible"] and q["source_sha256"] == h["source"]["sha256"] and q["core_sha256"] == h["core_sha256"], "QC/handoff mismatch")
    require(fingerprint(h["dictionary"]) == h["dictionary_sha256"] and core_hash(h) == h["core_sha256"], "handoff design/dictionary changed")
    require(h["scientific_confirmation"] and h["scientific_confirmation"]["core_sha256"] == h["core_sha256"], "confirmation binding missing")
    require(h["case_kind"] == q["case_kind"] and h["revision"] == q["revision"] and h["study_id"] == q["study_id"], "handoff identity mismatch")
    return {"status": "HANDOFF_RECEIVED", "study_id": h["study_id"], "revision": h["revision"], "case_kind": h["case_kind"], "statistics_run": False,
            "next": "按方案核查临床/统计适用性；合成演示不能转真实临床推断"}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("create"); p.add_argument("--spec", required=True); p.add_argument("--out", required=True)
    for action in ("resume", "confirm", "amend", "progress", "export", "qc"):
        p = sub.add_parser(action); p.add_argument("--package", required=True)
        if action in {"confirm", "progress"}: p.add_argument("--record", required=True)
        if action == "amend": p.add_argument("--changes", required=True)
        if action in {"export", "qc"}: p.add_argument("--out", required=True)
        if action == "export": p.add_argument("--rows", type=int, default=200); p.add_argument("--text-only", action="store_true")
        if action == "qc": p.add_argument("--workbook", required=True)
    p = sub.add_parser("accept-handoff"); p.add_argument("--handoff", required=True); p.add_argument("--source-override"); p.add_argument("--package-override")
    a = parser.parse_args(argv)
    try:
        if a.command == "create": result = create_package(read_json(a.spec), a.out)
        elif a.command == "resume": result = resume(a.package)
        elif a.command == "confirm": result = confirm_package(a.package, read_json(a.record))
        elif a.command == "amend": result = amend(a.package, read_json(a.changes))
        elif a.command == "progress": result = record_progress(a.package, read_json(a.record))
        elif a.command == "export": result = export_crf(a.package, a.out, a.rows, a.text_only)
        elif a.command == "qc": result = return_and_handoff(a.package, a.workbook, a.out)
        else: result = validate_handoff(a.handoff, a.source_override, a.package_override)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 2 if a.command == "qc" and not result["analysis_handoff_eligible"] else 0
    except (ValueError, KeyError, IndexError, TypeError, OSError, ImportError, StopIteration) as exc:
        print(json.dumps({"status": "NEEDS_INPUT_OR_CAPABILITY", "reason": str(exc), "next": "保留原工作包，按缺项续接，不自动安装或改原数据"}, ensure_ascii=False), file=sys.stderr)
        return 3


if __name__ == "__main__":
    sys.exit(main())

"""Ye1 local data/analysis executor. Stdlib orchestration; existing R does inference.

No installation, external API, generated executable code, or clinical schema defaults.
"""
from __future__ import annotations

# All data-bearing streams stay behind the controlled local adapter.
if __name__ == '__main__':
    import sys as _privacy_sys
    if '--help' not in _privacy_sys.argv[1:]:
        from privacy_boundary import cli_local_entry
        cli_local_entry()

import argparse, csv, hashlib, importlib.util, json, math, os, re, shutil, subprocess, sys, tempfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

VERSION = '0.1.0'
class Blocked(ValueError): pass
class EnvironmentBlocked(RuntimeError): pass
def now(): return datetime.now(timezone.utc).isoformat()
def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def digest(obj): return hashlib.sha256(json.dumps(obj, ensure_ascii=False, sort_keys=True, allow_nan=False).encode()).hexdigest()
def read_json(path): return json.loads(Path(path).read_text(encoding='utf-8'))
def write_json(path, obj): Path(path).write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False)+'\n', encoding='utf-8')
def new_out(path):
    p=Path(path)
    if p.exists(): raise Blocked('Output exists; preserve prior version and choose a new directory')
    p.mkdir(parents=True); return p
def require(cond, reason):
    if not cond: raise Blocked(reason)

def load_table(path, sheet=None):
    p=Path(path)
    if p.suffix.lower()=='.csv':
        with p.open(encoding='utf-8-sig', newline='') as f: rows=list(csv.reader(f))
    elif p.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        w=load_workbook(p, read_only=True, data_only=False)
        require(sheet in w.sheetnames if sheet else len(w.sheetnames)==1, 'Select the exact worksheet; do not guess')
        s=w[sheet] if sheet else w[w.sheetnames[0]]
        rows=[[v.isoformat() if isinstance(v,datetime) else ('' if v is None else str(v)) for v in r] for r in s.values]; w.close()
    else: raise Blocked('Only UTF-8 CSV or an explicitly selected XLSX sheet is supported')
    require(bool(rows), 'Empty table')
    header=rows[0]
    require(all(header) and len(header)==len(set(header)), 'Blank or duplicate column names; resolve header first')
    require(all(len(r)==len(header) for r in rows[1:]), 'Ragged rows; no silent truncation')
    return header, [dict(zip(header,r)) for r in rows[1:]]

def dictionary(spec, header):
    require(isinstance(spec.get('fields'),list) and spec['fields'], 'Dictionary requires fields')
    fields={f['name']:f for f in spec['fields']}
    require(len(fields)==len(spec['fields']), 'Duplicate dictionary fields')
    require(set(fields)<=set(header), 'Dictionary field absent from data: '+str(sorted(set(fields)-set(header))))
    require(spec.get('row_unit') and spec.get('key'), 'Confirm row_unit and record key; patient/encounter/visit are distinct')
    require(set(spec['key'])<=set(fields), 'Record key must be declared in dictionary')
    for f in fields.values():
        require(f.get('type') in {'number','integer','category','binary','date','string'}, 'Unsupported field type')
        if f['type'] in {'category','binary'}: require(f.get('levels') and len(set(f['levels']))==len(f['levels']), 'Declare unique category levels')
        if f['type']=='binary': require(len(f['levels'])==2, 'Binary field needs exactly two levels')
        if f.get('unit_column'): require(f['unit_column'] in header and f.get('unit'), 'Unit column/expected unit missing')
        if f.get('missing_reason_column'):require(f['missing_reason_column'] in header,'Missing-reason companion column absent')
        if f.get('pattern'):re.compile(f['pattern'])
    return fields

def missing(v, f): return v in f.get('missing_values',[''])
def parse(v,f):
    if missing(v,f): return None
    if v.startswith('='): raise ValueError('formula needs verified values; cached formula not trusted')
    kind=f['type']
    if kind in {'number','integer'}:
        n=float(v)
        if not math.isfinite(n) or (kind=='integer' and not n.is_integer()): raise ValueError('invalid finite number/integer')
        return n
    if kind in {'category','binary'} and v not in f['levels']: raise ValueError('undeclared category')
    if kind=='date': return datetime.fromisoformat(v).isoformat()
    if f.get('pattern') and not re.fullmatch(f['pattern'],v):raise ValueError('declared text pattern mismatch')
    return v

def qc(header, rows, spec):
    fields=dictionary(spec,header); issues=[]; counts={}; parsed=[]
    def issue(code, row=None, field=None, severity='error', detail=None):
        issues.append({'code':code,'row':row,'field':field,'severity':severity,'detail':detail})
    for i,row in enumerate(rows,2):
        vals={}
        for name,f in fields.items():
            try: val=parse(row[name],f)
            except (ValueError,OverflowError) as e: val=None; issue('INVALID_VALUE',i,name,detail='Local parse failure; inspect protected QC input')
            vals[name]=val
            if missing(row[name],f) and f.get('required'): issue('REQUIRED_MISSING',i,name)
            if val is not None and f['type'] in {'number','integer'}:
                for bound,op in [('min',lambda a,b:a<b),('max',lambda a,b:a>b)]:
                    if bound in f and op(val,f[bound]): issue('OUTSIDE_DECLARED_RANGE',i,name,'warning','Review; never delete automatically')
            if val is not None and f.get('unit_column') and row[f['unit_column']]!=f['unit']:
                issue('UNIT_MISMATCH',i,name,detail='Expected '+f['unit'])
            if f.get('missing_reason_column'):
                reason=row[f['missing_reason_column']]
                if val is not None and reason:issue('VALUE_AND_MISSING_REASON',i,name)
                if missing(row[name],f) and not reason:issue('MISSING_REASON_REQUIRED',i,name)
        for rule in spec.get('time_rules',[]):
            a,b=rule['before'],rule['after']
            require(a in fields and b in fields and fields[a]['type']==fields[b]['type']=='date','Time rule requires declared date fields')
            if vals.get(a) is not None and vals.get(b) is not None:
                try: bad=datetime.fromisoformat(vals[a])>datetime.fromisoformat(vals[b])
                except TypeError: issue('TIMEZONE_MISMATCH',i,a); bad=False
                if bad: issue('TIME_ORDER',i,a,detail=a+' must not follow '+b)
        parsed.append(vals)
    keys=[tuple(row[k] for k in spec['key']) for row in rows]
    for key,n in Counter(keys).items():
        if n>1: issue('DUPLICATE_KEY',detail={'count':n})
    if any(any(missing(row[k],fields[k]) for k in spec['key']) for row in rows): issue('MISSING_KEY')
    exact=sum(n-1 for n in Counter(tuple(row[h] for h in header) for row in rows).values() if n>1)
    if exact: issue('DUPLICATE_ROW',detail={'excess_rows':exact})
    subject=spec.get('subject_id')
    if subject:
        require(subject in fields, 'subject_id not in dictionary')
        if len(set(r[subject] for r in rows))<len(rows): issue('REPEATED_SUBJECT',severity='warning',detail='Independent-row inference blocked; use appropriate repeated/clustered method')
    for name,f in fields.items():
        nmiss=sum(missing(r[name],f) for r in rows)
        counts[name]={'observed':len(rows)-nmiss,'missing':nmiss,'missing_fraction':nmiss/len(rows) if rows else None,'unit':f.get('unit'),'type':f['type']}
    return {'status':'ISSUES' if issues else 'READY','n_rows':len(rows),'row_unit':spec['row_unit'],'key':spec['key'],'columns':counts,'undeclared_columns':sorted(set(header)-set(fields)),'issues':issues,'clinical_validity':'HOST_REQUIRED'}, parsed

def import_crf(crf, table_name, out, data=None, subject_id=None):
    """Explicit field adapter for the current project CRF dictionary.tables contract.
    Does not replace CRF accept-handoff or imply research/collection authorization.
    """
    source=read_json(crf);tables=source.get('dictionary',source).get('tables',[])
    matches=[t for t in tables if t['name']==table_name];require(len(matches)==1,'Select one exact CRF table')
    table=matches[0];fields=[];header_map={};pending=[]
    for old in table['fields']:
        f={**old,'type':'date' if old['type']=='datetime' else old['type'],'missing_values':[''],'crf_required':old.get('required',False)}
        if old['type']=='category':f['levels']=list(old['allowed_values'])
        if old.get('missing_codes'):f['required']=False;f['missing_reason_column']=f['name']+'__missing'
        fields.append(f)
        title=f['name']+' | '+old['label']+(' ('+old['unit']+')' if old['unit']!='无' else '')+(' *' if old.get('required') else '')
        header_map[title]=f['name']
        if old.get('missing_codes'):
            name=f['name']+'__missing';fields.append({'name':name,'type':'category','levels':list(old['missing_codes']),'label':'缺失原因','definitions':old['missing_codes']});header_map[name+' | 缺失原因']=name
    time_rules=[]
    by={f['name']:f for f in fields}
    for r in table.get('constraints',[]):
        if r['kind']=='order' and r.get('op') in {'<=','>='} and by[r['left']]['type']==by[r['right']]['type']=='date':
            time_rules.append({'before':r['left'] if r['op']=='<=' else r['right'],'after':r['right'] if r['op']=='<=' else r['left']})
        else:pending.append(r)
    if table.get('foreign_keys'):pending.append({'foreign_keys':table['foreign_keys'],'action':'Validate related tables in CRF module; no automatic joins'})
    spec={'row_unit':table['grain'],'key':table['primary_key'],'fields':fields,'time_rules':time_rules,'crf_source':{'sha256':sha(crf),'table':table_name,'source_schema':source.get('schema_version','dictionary.tables')},'pending_crf_rule_review':pending}
    if subject_id:require(subject_id in by,'subject_id absent from CRF');spec['subject_id']=subject_id
    target=new_out(out);write_json(target/'dictionary.json',spec);write_json(target/'header-map.json',header_map);shutil.copyfile(crf,target/'crf-dictionary-source.json')
    result={'status':'REVIEW_PENDING','dictionary_sha256':digest(spec),'eligibility':'CRF accept-handoff and downstream scientific checks remain required','pending_rules':pending}
    if data:
        header,rows=load_table(data,table_name if Path(data).suffix.lower()=='.xlsx' else None)
        # CRF final delivery uses readable Chinese first, newline stable ID.
        # Retain the exact legacy schema for existing externally supplied workbooks.
        current_map={}
        for old_field in table['fields']:
            title=old_field['label']+(' ('+old_field['unit']+')' if old_field['unit']!='无' else '')+(' *' if old_field.get('required') else '')+'\n'+old_field['name']
            current_map[title]=old_field['name']
            if old_field.get('missing_codes'):current_map['缺失原因\n'+old_field['name']+'__missing']=old_field['name']+'__missing'
        if header==list(current_map):
            header_map=current_map
            result['header_contract']='CRF_CHINESE_NEWLINE_ID'
        else:
            require(header==list(header_map),'CRF header changed; no inferred mapping')
            result['header_contract']='CRF_LEGACY_ID_PIPE_LABEL'
        write_json(target/'header-map.json',header_map)
        normalized=[];row_map=[]
        for i,row in enumerate(rows,2):
            if all(v=='' for v in row.values()):continue
            normalized.append({header_map[k]:v for k,v in row.items()});row_map.append({'normalized_row':len(normalized)+1,'source_row':i})
        save_csv(target/'normalized.csv',list(header_map.values()),normalized);shutil.copyfile(data,target/('source'+Path(data).suffix));write_json(target/'source-row-map.json',row_map)
        report,_=qc(list(header_map.values()),normalized,spec);write_json(target/'qc.json',report)
        result.update(source_sha256=sha(data),normalized_sha256=sha(target/'normalized.csv'),blank_rows_skipped=len(rows)-len(normalized),normalized_rows=len(normalized))
    write_json(target/'import-receipt.json',result);return result

def save_csv(path,header,rows):
    with Path(path).open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=header);w.writeheader();w.writerows(rows)

def clean(data, spec, rules, out, sheet=None):
    header,rows=load_table(data,sheet); dictionary(spec,header); original=[dict(r) for r in rows]; changes=[]
    require(rules.get('source_sha256')==sha(data),'Cleaning rules must bind exact source SHA256')
    require(rules.get('reason') and rules.get('confirmed_by') and rules.get('confirmed_at'), 'Cleaning rules require explicit provenance/confirmation')
    for rule in rules.get('rules',[]):
        op=rule.get('op'); require(op in {'map','convert_unit','set_missing','correct_cell','drop_rows'},'Unsupported cleaning rule')
        if op=='drop_rows':
            indices=rule['source_rows'];require(rule.get('reason') and all(isinstance(x,int) and 2<=x<len(original)+2 for x in indices),'Explicit source rows/reason required for deletion')
            for r in indices:
                require(rows[r-2] is not None,'Row already removed');changes.append({'source_row':r,'op':op,'before':rows[r-2],'reason':rule['reason']});rows[r-2]=None
            continue
        field=rule['field'];require(field in header,'Unknown cleaning field')
        if op=='correct_cell':
            require(isinstance(rule.get('source_row'),int) and 2<=rule['source_row']<len(original)+2,'Correction source row outside source data')
            require(rows[rule['source_row']-2] is not None,'Correction targets a removed row')
        if op=='convert_unit':require(rule.get('unit_column') in header and rule.get('from_unit') and rule.get('to_unit'),'Explicit unit column/from/to required')
        for i,row in enumerate(rows,2):
            if row is None: continue
            old=row[field]; new=old
            if op=='map': new=rule['mapping'].get(old,old)
            if op=='set_missing' and old in rule['values']: new=''
            if op=='correct_cell' and i==rule['source_row']:
                require(old==rule['expected_old'],'Correction old value mismatch'); new=str(rule['value'])
            if op=='convert_unit' and row[rule['unit_column']]==rule['from_unit'] and old!='':
                number=float(old)*float(rule['factor'])+float(rule.get('offset',0));require(math.isfinite(number),'Nonfinite converted value')
                new=format(number,'.15g');uc=rule['unit_column']
                changes.append({'source_row':i,'field':uc,'before':row[uc],'after':rule['to_unit'],'rule':rule});row[uc]=rule['to_unit']
            if new!=old: changes.append({'source_row':i,'field':field,'before':old,'after':new,'rule':rule});row[field]=new
    target=new_out(out);shutil.copyfile(data,target/('source'+Path(data).suffix));save_csv(target/'cleaned.csv',header,[r for r in rows if r is not None]);write_json(target/'rules.json',rules)
    report,_=qc(header,[r for r in rows if r is not None],spec);write_json(target/'qc-after.json',report)
    receipt={'status':'REVIEW_PENDING','source_sha256':sha(data),'cleaned_sha256':sha(target/'cleaned.csv'),'changes':changes,'rows_before':len(original),'rows_after':sum(r is not None for r in rows),'outliers_removed_automatically':False};write_json(target/'changes.json',receipt);return receipt

def doctor(rscript=None):
    rscript=rscript or os.environ.get('YE1_RSCRIPT') or shutil.which('Rscript')
    r={'python':sys.version,'xlsx_available':bool(importlib.util.find_spec('openpyxl')),'rscript':rscript,'text_host_possible':True,'installs':False}
    if rscript:
        try:
            p=subprocess.run([rscript,'--vanilla','-e','stopifnot(requireNamespace("jsonlite",quietly=TRUE)); cat(R.version.string,"; jsonlite ",as.character(packageVersion("jsonlite")))'],capture_output=True,text=True,timeout=30)
            r.update(status='READY' if p.returncode==0 else 'BLOCKED_ENVIRONMENT',details=p.stdout+p.stderr)
        except (OSError,subprocess.SubprocessError) as e:r.update(status='BLOCKED_ENVIRONMENT',details=str(e))
    else:r.update(status='BLOCKED_ENVIRONMENT',details='Existing Rscript + jsonlite unavailable')
    r['fallback']='Host can plan/interpret existing results. No simulated statistics; QC/clean CSV use Python stdlib.';return r

def run_r(payload, target, rscript=None):
    env=doctor(rscript)
    if env['status']!='READY':raise EnvironmentBlocked(env['details'])
    write_json(target/'worker-input.json',payload)
    p=subprocess.run([env['rscript'],'--vanilla',str(Path(__file__).with_name('stats_worker.R')),str(target/'worker-input.json'),str(target/'worker-output.json')],capture_output=True,text=True,timeout=180)
    (target/'worker.log').write_text(p.stdout+p.stderr,encoding='utf-8')
    if p.returncode!=0 or not (target/'worker-output.json').is_file():raise EnvironmentBlocked('R execution failed; retain logs, no statistics inferred')
    result=read_json(target/'worker-output.json');return result,env

def analysis_dependencies(task,spec):
    names=set(task.get('variables',[])+task.get('predictors',[]))
    for key in ('outcome','group','paired_with'):
        if task.get(key):names.add(task[key])
    return sorted(names|set(spec['key'])|({spec['subject_id']} if spec.get('subject_id') else set()))

def check_plan(plan, spec, header, report):
    from confirmation_evidence import validate as validate_confirmation
    validate_confirmation(plan,'confirmation_statement')
    require(plan.get('design')=='independent' and plan.get('row_unit')==spec['row_unit'],'Only explicit independent-row design supported; clarify row unit/longitudinal/clustered data')
    require(plan.get('research_question') and plan.get('population') and plan.get('estimand'),'Plan needs question, population and estimand')
    require(plan.get('confirmed_by') and plan.get('confirmed_at') and plan.get('confirmation_statement'),'Record actual method/definition confirmation, never fabricate')
    if spec.get('pending_crf_rule_review'):require(plan.get('crf_rule_review',{}).get('dictionary_sha256')==digest(spec) and plan['crf_rule_review'].get('statement'),'Resolve CRF constraints/foreign-key review before analysis')
    ids=[t['id'] for t in plan['tasks']];require(len(set(ids))==len(ids),'Duplicate analysis IDs')
    fields=dictionary(spec,header)
    for t in plan['tasks']:
        require(t['kind'] in {'describe','compare','linear','logistic','prediction'},'Unsupported analysis route')
        require(set(analysis_dependencies(t,spec))<=set(fields),'Analysis uses undeclared fields')
        require(t.get('missing') in {'complete_case','available_case','train_median'},'Declare missing-data handling')
        if t['kind']=='describe':require(t.get('missing')=='available_case','Descriptive executor is explicitly variable-wise available-case')
        if t['kind']!='describe':
            require(t.get('missing')!='available_case','Inferential analyses require one explicit joint sample')
            require(plan.get('independence_basis'),'Document independence basis')
        if t['kind'] in {'linear','logistic'}:require(t.get('covariate_rationale') and t.get('assumptions_review'),'Declare confounding/covariate rationale and assumptions')
        if t['kind']=='prediction':
            require(t.get('feature_review')=='CONFIRMED' and t.get('prediction_time') and t.get('outcome_horizon'),'Review feature timing and leakage before prediction')
            for n in t['predictors']:
                require(n!=t['outcome'] and n not in spec['key'] and fields[n].get('role') not in {'identifier','outcome'},'Identifier/outcome leakage forbidden')
                require(fields[n].get('available_at')=='prediction_time' and not fields[n].get('derived_from_outcome'),'Feature not known at prediction time: '+n)
        if t.get('missing')=='train_median':require(t['kind']=='prediction','Train-only median imputation applies to prediction only')

def export_tables_and_plot_sources(target,results,plan,spec,parsed,data_sha):
    table=[];plots=[];tasks={t['id']:t for t in plan['tasks']};fields={f['name']:f for f in spec['fields']}
    for result in results:
        if result['status']!='OK':continue
        task=tasks[result['id']]
        if result['kind']=='describe':
            for s in result['statistics']:
                base={'analysis_id':result['id'],'group':s['group'],'variable':s['variable'],'unit':s['unit'],'n_total':s['n_total'],'n_observed':s['n_observed'],'n_missing':s['n_missing']}
                if 'levels' in s:
                    for item in s['levels']:
                        table.append({**base,'statistic':'count','category':item['level'],'value':item['count']});table.append({**base,'statistic':'percent_observed','category':item['level'],'value':item['percent']})
                else:
                    for statistic in ('mean','sd'):table.append({**base,'statistic':statistic,'category':'','value':s[statistic]})
                    for statistic,value in (s.get('quantiles') or {}).items():table.append({**base,'statistic':statistic,'category':'','value':value})
            for n in task['variables']:
                plots.append({'analysis_id':result['id'],'kind':'distribution' if fields[n]['type'] in {'number','integer'} else 'category_counts','variable':n,'unit':fields[n].get('unit'),'group_variable':task.get('group'),'rows':[{'source_row':i,'value':r[n],'group':r.get(task.get('group')) if task.get('group') else 'All'} for i,r in enumerate(parsed,2)],'missing_handling':'Retain null values; graph module must disclose n plotted/missing and exclude missing groups explicitly'})
        elif result['kind'] in {'linear','logistic'}:
            plots.append({'analysis_id':result['id'],'kind':'forest','outcome':task['outcome'],'event':task.get('event'),'coefficients':result['coefficients'],'n_used':result['n_used'],'scale':'odds_ratio' if result['kind']=='logistic' else 'linear_coefficient','include_intercept':False})
    if table:save_csv(target/'descriptive-table.csv',['analysis_id','group','variable','statistic','category','value','unit','n_total','n_observed','n_missing'],table)
    write_json(target/'statistical-graphics-handoff.json',{'schema':'ye1-statistical-plots-v1','data_sha256':data_sha,'items':plots,'rendered':False,'theme_owner':'Main integration graphics module'})

def analyze(data, spec, plan, out, sheet=None, previous=None, confirmation=None, rscript=None):
    header,rows=load_table(data,sheet);report,parsed=qc(header,rows,spec);check_plan(plan,spec,header,report)
    require(not any(x['severity']=='error' for x in report['issues']),'QC errors unresolved; run qc/clean before inference')
    require(not any(x['code']=='REPEATED_SUBJECT' for x in report['issues']),'Repeated subjects require a different method; independent-row engine blocked')
    ranges=[x for x in report['issues'] if x['code']=='OUTSIDE_DECLARED_RANGE']
    if ranges:require(plan.get('range_review',{}).get('source_sha256')==sha(data) and plan['range_review'].get('statement'),'Review out-of-range values on this exact data; never silently trim')
    prior=read_json(Path(previous)/'analysis.json') if previous else None
    if prior:
        for name,h in prior['frozen_files'].items(): require(sha(Path(previous)/name)==h,'Prior baseline changed: '+name)
        changed_definition=prior['plan_sha256']!=digest(plan) or prior['dictionary_sha256']!=digest(spec)
        if changed_definition:
            from confirmation_evidence import validate as validate_confirmation
            validate_confirmation(confirmation or {},'statement')
            require(confirmation and confirmation.get('from_plan_sha256')==prior['plan_sha256'] and confirmation.get('to_plan_sha256')==digest(plan) and confirmation.get('to_dictionary_sha256')==digest(spec) and all(confirmation.get(k) for k in ('actor','date','statement')), 'Major definition/method/dictionary change requires explicit version-bound confirmation')
    else:changed_definition=False
    target=new_out(out);shutil.copyfile(data,target/('source'+Path(data).suffix));write_json(target/'dictionary.json',spec);write_json(target/'plan.json',plan);write_json(target/'qc.json',report)
    signatures={t['id']:digest({'task':t,'definition':{k:v for k,v in plan.items() if k!='tasks'},'fields':{n:next(f for f in spec['fields'] if f['name']==n) for n in analysis_dependencies(t,spec)},'data':[[r[n] for n in analysis_dependencies(t,spec)] for r in rows]}) for t in plan['tasks']}
    old_by={r['id']:r for r in prior['results']} if prior else {};results=[];execute=[];reused=[]
    for t in plan['tasks']:
        if prior and not changed_definition and prior['dependency_signatures'].get(t['id'])==signatures[t['id']]: results.append(old_by[t['id']]);reused.append(t['id'])
        else:execute.append(t)
    env=None
    if execute:
        worker,env=run_r({'rows':parsed,'dictionary':spec,'tasks':execute,'data_sha256':sha(data)},target,rscript);results.extend(worker['results'])
    order={t['id']:i for i,t in enumerate(plan['tasks'])};results.sort(key=lambda x:order[x['id']])
    changed=not prior or prior['data_sha256']!=sha(data) or changed_definition
    changed_ids=[t['id'] for t in execute];removed=sorted(set(old_by)-set(signatures))
    old_tasks=read_json(Path(previous)/'plan.json')['tasks'] if prior else []
    downstream=sorted({a for t in plan['tasks'] if t['id'] in changed_ids for a in t.get('downstream',[])} | {a for t in old_tasks if t['id'] in removed for a in t.get('downstream',[])})
    record={'schema':'ye1-analysis-v1','engine_version':VERSION,'created_at':now(),'status':'PARTIAL' if any(r['status']!='OK' for r in results) else 'COMPUTED_REVIEW_PENDING','data_sha256':sha(data),'plan_sha256':digest(plan),'dictionary_sha256':digest(spec),'previous':str(Path(previous).resolve()) if previous else None,'previous_analysis_sha256':sha(Path(previous)/'analysis.json') if prior else None,'results':results,'dependency_signatures':signatures,'recomputed':changed_ids,'reused':reused,'removed':removed,'affected_downstream':downstream,'downstream_other':'Host must check unregistered consumers; no full manuscript rewrite required','confirmation_status':'STALE' if prior and changed else 'PENDING','old_confirmations_valid':False if changed else None,'environment':env,'definition_change':changed_definition,'definition_confirmation':confirmation,'confirmation_source_kind':plan.get('source_kind','unknown'),'frozen_files':{p.name:sha(p) for p in target.iterdir() if p.is_file()},'semantic_review':'HOST_REQUIRED; no causal or significance-language automation'}
    export_tables_and_plot_sources(target,results,plan,spec,parsed,sha(data))
    graphics=[]
    for r in results:
        if r.get('plot_payload'):graphics.append({'analysis_id':r['id'],'data_sha256':sha(data),'validation':r.get('validation'),**r['plot_payload']})
    write_json(target/'graphics-handoff.json',{'schema':'ye1-prediction-plots-v1','items':graphics,'dca':'Not generated; requires clinically justified threshold range, decision setting and adequate validation'})
    record['frozen_files']={p.name:sha(p) for p in target.iterdir() if p.is_file()}
    write_json(target/'analysis.json',record)
    write_json(target/'state-patch.json',{'schema':'ye1-state-patch-v1','active_versions':{'data':sha(data),'analysis':sha(target/'analysis.json'),'plan':digest(plan),'dictionary':digest(spec)},'facts':[{'status':'CONFIRMED','source':str(target/'analysis.json'),'meaning':'Machine computation occurred; scientific acceptance still pending','date':now()}],'pending':['Review method assumptions, results and downstream scope'],'invalidate_manuscript_confirmation':bool(prior and changed),'invalidate_results':changed_ids+removed,'preserve_unaffected_results':reused,'merge_instruction':'Host merge only these owned keys into existing doctor-paper-state.json; preserve unrelated facts/references/terminology and retain prior state'})
    return record

def sample_size(spec,out,rscript=None):
    target=new_out(out);result,env=run_r({'sample_size':spec},target,rscript);result['environment']=env;write_json(target/'sample-size.json',result);return result

def main():
    p=argparse.ArgumentParser(description=__doc__);s=p.add_subparsers(dest='cmd',required=True)
    d=s.add_parser('doctor');d.add_argument('--rscript')
    for mode in ('qc','dictionary','clean','analyze'):
        a=s.add_parser(mode);a.add_argument('data');a.add_argument('--dictionary',required=True);a.add_argument('--sheet');a.add_argument('--out',required=True)
        if mode=='clean':a.add_argument('--rules',required=True)
        if mode=='analyze':
            a.add_argument('--plan',required=True);a.add_argument('--previous');a.add_argument('--confirmation');a.add_argument('--rscript')
    a=s.add_parser('sample-size');a.add_argument('spec');a.add_argument('--out',required=True);a.add_argument('--rscript')
    a=s.add_parser('import-crf');a.add_argument('crf');a.add_argument('--table',required=True);a.add_argument('--data');a.add_argument('--subject-id');a.add_argument('--out',required=True)
    x=p.parse_args()
    if x.cmd=='doctor':r=doctor(x.rscript)
    elif x.cmd=='import-crf':r=import_crf(x.crf,x.table,x.out,x.data,x.subject_id)
    elif x.cmd=='sample-size':r=sample_size(read_json(x.spec),x.out,x.rscript)
    else:
        spec=read_json(x.dictionary)
        if x.cmd=='clean':r=clean(x.data,spec,read_json(x.rules),x.out,x.sheet)
        elif x.cmd=='analyze':r=analyze(x.data,spec,read_json(x.plan),x.out,x.sheet,x.previous,read_json(x.confirmation) if x.confirmation else None,x.rscript)
        else:
            h,rows=load_table(x.data,x.sheet);r,_=qc(h,rows,spec);target=new_out(x.out);write_json(target/'qc.json',r);write_json(target/'dictionary.json',spec)
            if x.cmd=='dictionary':
                write_json(target/'crf-fields.json',{'schema':'ye1-crf-fields-v1','row_unit':spec['row_unit'],'key':spec['key'],'fields':spec['fields'],'mapping_status':'Explicit field-name adapter required if CRF uses another schema; no inferred clinical semantics'})
    print(json.dumps({'status':r.get('status'),'output_retained_locally':bool(getattr(x,'out',None)),'detail':None},ensure_ascii=False))
    return 3 if r.get('status')=='BLOCKED_ENVIRONMENT' else 0

if __name__=='__main__':
    try:sys.exit(main())
    except (Blocked,KeyError,TypeError,ValueError) as e:print(json.dumps({'status':'BLOCKED_VALIDATION','reason':'Local validation failed; inspect protected inputs'},ensure_ascii=False),file=sys.stderr);sys.exit(2)
    except (EnvironmentBlocked,ImportError,OSError,subprocess.SubprocessError) as e:print(json.dumps({'status':'BLOCKED_ENVIRONMENT','reason':'Local validation failed; inspect protected inputs','next_step':'Retain existing results/text; select existing runtime. Never install automatically.'},ensure_ascii=False),file=sys.stderr);sys.exit(3)

"""Local workbook/CSV structural profile, independent of clinical topic/schema.
No cell examples or identifiers are emitted. The host reads relevant raw cells locally.
"""

# All data-bearing streams stay behind the controlled local adapter.
if __name__ == '__main__':
    import sys as _privacy_sys
    if '--help' not in _privacy_sys.argv[1:]:
        from privacy_boundary import cli_local_entry
        cli_local_entry()

import argparse,csv,json
from collections import Counter
from pathlib import Path

def profile(path):
    p=Path(path);sheets=[]
    if p.suffix.lower()=='.csv':
        with p.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.reader(f))
        sources=[('CSV',rows,[])]
    elif p.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        w=load_workbook(p,data_only=False,read_only=False)
        sources=[(s.title,[[x.value for x in row] for row in s.iter_rows()],[str(x) for x in s.merged_cells.ranges]) for s in w]
    else: raise ValueError('Supported: local .xlsx or UTF-8 CSV only.')
    for name,rows,merged in sources:
        # Propose rows, never assert a clinical header or join key automatically.
        candidates=[]
        for i,row in enumerate(rows[:10]):
            vals=[x for x in row if x is not None and str(x).strip()]
            if len(vals)>1 and all(isinstance(x,str) and not x.startswith('=') for x in vals):
                candidates.append({'row':i+1,'nonempty_cells':len(vals)})
        sheets.append({'sheet':name,'physical_rows':len(rows),'max_columns':max(map(len,rows),default=0),'merged_cells':merged,'header_row_candidates':candidates,'formula_cells':sum(isinstance(x,str) and x.startswith('=') for row in rows for x in row),'blank_cells':sum(x is None or str(x).strip()=='' for row in rows for x in row)})
    return {'format':p.suffix.lower(),'sheets':sheets,'scope':'Structural suggestions only; host must confirm row meaning, relations, units, timing and scientific goal. No fixed CRP schema required.'}

if __name__=='__main__':
    a=argparse.ArgumentParser(description=__doc__);a.add_argument('input');a.add_argument('--out',required=True);x=a.parse_args();p=Path(x.out)
    if p.exists(): raise SystemExit('Output exists; choose a new file.')
    p.write_text(json.dumps(profile(x.input),ensure_ascii=False,indent=2)+'\n',encoding='utf8'); print('Profile written locally: '+str(p))

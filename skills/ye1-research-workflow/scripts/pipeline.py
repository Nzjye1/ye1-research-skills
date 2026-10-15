#!/usr/bin/env python3
"""Local bounded descriptive Excel/CSV -> bilingual manuscript, no model/network calls."""
from __future__ import annotations

# All data-bearing streams stay behind the controlled local adapter.
if __name__ == '__main__':
    import sys as _privacy_sys
    if '--help' not in _privacy_sys.argv[1:]:
        from privacy_boundary import cli_local_entry
        cli_local_entry()

import argparse, csv, hashlib, json, math, re, statistics as st, sys
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from collections import Counter
import materialize
import environment
import subprocess
from gate_contract import require_gate
from graphics_bridge import plots as v02plots

class Blocked(ValueError): pass

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(p,x): Path(p).write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf8')
def load(p): return json.loads(Path(p).read_text(encoding='utf8'))
def fmt(x): return 'NA' if x is None else f'{x:.1f}'
def norm(x):
    if x is None: return None
    s=str(x).strip()
    return None if s.lower() in {'','na','nan','n/a','未知','未查','—','-'} else s

def read_table(path,c):
    p=Path(path)
    if p.suffix.lower()=='.csv':
        with p.open(encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f)), {'merged_cells':[], 'sheet':'CSV'}
    if p.suffix.lower()!='.xlsx': raise Blocked('仅支持本地 .xlsx / UTF-8 CSV；请另存副本，不覆盖原表。')
    from openpyxl import load_workbook
    w=load_workbook(p,data_only=False,read_only=False)
    sheet=c.get('sheet')
    if not sheet and len(w.sheetnames)!=1: raise Blocked('工作簿含多表，请先确认分析表及表间关系；不能自动拼接。')
    s=w[sheet or w.sheetnames[0]]; header=c.get('header_row',1)
    cols=[norm(x.value) for x in s[header]]
    if None in cols or len(set(cols))!=len(cols): raise Blocked('表头有空白或重名，请确认 header_row 或提供字段映射。')
    rows=[]
    for line in s.iter_rows(min_row=header+1):
        if any(x.data_type=='f' for x in line): raise Blocked('发现公式：请在副本中转为经核实的值；不读取可能过时的缓存。')
        vals=[x.value.isoformat() if isinstance(x.value,(date,datetime)) else x.value for x in line]
        if any(norm(x) is not None for x in vals): rows.append(dict(zip(cols,vals)))
    return rows, {'merged_cells':[str(x) for x in s.merged_cells.ranges], 'sheet':s.title}

def inspect(path,c,out):
    rows,structure=read_table(path,c); cols=list(rows[0]) if rows else []
    exact=len(rows)-len({json.dumps(r,sort_keys=True,ensure_ascii=False,default=str) for r in rows})
    out=Path(out)
    if out.exists(): raise Blocked('输出目录已存在；用新目录保留历史。')
    out.mkdir(parents=True)
    save(out/'intake.json',{'columns':cols,'rows':len(rows),'exact_duplicate_rows':exact,'structure':structure,'source_sha256':sha(path),'status':'FEASIBILITY_ONLY'})
    follow=c.get('followup',{}); closed=follow.get('status')=='unavailable' and follow.get('can_supplement') is False
    text=f'''# 先看看这份表能做什么\n\n读到 {len(rows)} 行、{len(cols)} 列，发现 {exact} 行完全重复。原表未修改。\n\n现有资料可以先做变量分布、缺失情况；确认每行代表谁、编号关系、单位和时点后，可按人或首次住院描述。是否能做关联或预后，要看结局定义和时间是否够用，不能根据列名猜。\n\n{'已记录：随访没有，也补不到。本轮不再追问随访，不能将其当作无事件；只研究现有住院资料。' if closed else '随访是否存在、是否能从病历补齐尚未确认；如果目标需要随访，集中确认一次即可。'}\n\n必要澄清：研究目标；一行代表人还是住院；编号如何关联；检验单位及首次/末次时点；允许的清理规则。\n现实可补：原病历可核的单位、日期、结局编码，具体能否补及成本由医生确认。补不到就限制对应结论。\n可选增强：有真实时间和结局后再另定关联/预后方案；不承诺补后显著。\n\n这里只做可行性反馈，不出图、不计算研究结论。字段：{', '.join(cols)}。\n'''
    (out/'医生首页.md').write_text(text,encoding='utf8')

def decimal_value(v,kind,unit):
    s=norm(v)
    if s is None: return None
    s=s.replace('，',',').replace(',','')
    m=re.fullmatch(r'([+-]?(?:\d+(?:\.\d*)?|\.\d+))\s*(.*)',s)
    if not m: raise Blocked(f'无法确认数值/单位：{v!r}；请核实，不静默删掉。')
    n=Decimal(m[1]); suffix=m[2].strip()
    if suffix and suffix!=unit and not(kind=='crp' and suffix in {'mg/L','mg/dL'}): raise Blocked(f'未知单位 {suffix}，目标 {unit}；需要科学确认。')
    if kind=='crp' and (suffix or unit)=='mg/dL': n*=10
    if n<0: raise Blocked(f'{kind} 出现负值：{v}，请确认。')
    if kind=='age' and n>120: raise Blocked('年龄超出120岁；请核实，不能自动修正。')
    return n

def summary(vals):
    v=sorted(float(x) for x in vals if x is not None)
    def q(p):
        if not v: return None
        z=(len(v)-1)*p; i=int(z); return v[i]+(v[min(i+1,len(v)-1)]-v[i])*(z-i)
    return {'n':len(v),'mean':st.mean(v) if v else None,'median':q(.5),'q1':q(.25),'q3':q(.75),'min':min(v) if v else None,'max':max(v) if v else None}

def scientific_contract(c):
    return {k:c.get(k) for k in ['mapping','units','scope','paired_definition','followup','objective','population','method']}

def clean(path,c):
    units=c.get('units',{})
    if units.get('age')!='years' or units.get('stay_days')!='days' or any(units.get(k) not in {'mg/L','mg/dL'} for k in ['crp_first','crp_last']):
        raise Blocked('固定分析量纲应为年龄years、日数days、CRP mg/L或mg/dL；不能重新命名临床指标来迁就模式。')
    if c.get('scope')!='first_admission_per_person' or c.get('method')!='descriptive-v1': raise Blocked('当前自动引擎只支持每人首次住院描述。新方法需单独确认及实现。')
    if c.get('paired_definition')!='last_recorded_minus_first_recorded': raise Blocked('需要明确首次/末次的含义；不能默认治疗前后。')
    confirm=c.get('confirmation',{})
    if confirm.get('status')!='confirmed' or not confirm.get('source') or not confirm.get('date'): raise Blocked('分析方案尚未确认：请集中确认层级、单位、时点及描述范围。')
    if confirm.get('kind')=='simulated' and c.get('synthetic') is not True: raise Blocked('模拟确认只能用于明确标记的合成输入。')
    if confirm.get('kind') not in {'simulated','doctor'}: raise Blocked('确认来源须为 doctor 或 simulated。')
    rows,structure=read_table(path,c); mp=c['mapping']
    required={'person','admission','admit_date','age','crp_first','crp_last','oxygen','stay_days'}
    if not required<=mp.keys(): raise Blocked('当前描述模式缺少字段映射：'+str(sorted(required-mp.keys())))
    if rows and not set(mp.values())<=rows[0].keys(): raise Blocked('字段映射含不存在的列。')
    seen=set(); allrows=[]; duplicate=0
    for r in rows:
        key=json.dumps(r,sort_keys=True,ensure_ascii=False,default=str)
        if key in seen: duplicate+=1; continue
        seen.add(key)
        p=norm(r[mp['person']]); a=norm(r[mp['admission']])
        if not p or not a: raise Blocked('人/住院编号缺失：需确认记录归属。')
        try: dt=datetime.fromisoformat(str(r[mp['admit_date']]).strip()).date()
        except Exception: raise Blocked('入院日期不是明确ISO日期；请确认日期格式。')
        x={'person':p,'admission':a,'admit_date':dt.isoformat()}
        for k in ['age','crp_first','crp_last','stay_days']:
            x[k]=decimal_value(r[mp[k]],'crp' if k.startswith('crp') else ('age' if k=='age' else 'days'),c['units'][k])
        o=norm(r[mp['oxygen']]); values=c.get('oxygen_values',{'是':1,'否':0,'1':1,'0':0})
        if o is not None and o not in values: raise Blocked('氧疗编码不明；请确认 '+str(o))
        x['oxygen']=values.get(o)
        if x['oxygen'] not in {0,1,None}: raise Blocked('氧疗映射必须是0/1。')
        x['death_label']=norm(r.get(mp.get('death_label',''))) if mp.get('death_label') else None
        allrows.append(x)
    if len({x['admission'] for x in allrows})!=len(allrows): raise Blocked('同一住院编号有不完全相同的记录；需核重复/冲突，不能任选一条。')
    first={}; alerts=[]
    for x in sorted(allrows,key=lambda r:(r['admit_date'],r['admission'])):
        if x['person'] in first and x['admit_date']==first[x['person']]['admit_date']: raise Blocked('同一人同日多次住院，首次定义不明确。')
        first.setdefault(x['person'],x)
    for p in first:
        records=sorted([x for x in allrows if x['person']==p],key=lambda x:x['admit_date'])
        for i,x in enumerate(records[:-1]):
            if x['death_label'] in {'是','1','死亡'}: alerts.append({'person':p,'issue':'表载死亡后还有入院；结局未核，不进行死亡/生存分析。'})
    selected=list(first.values())
    return selected,allrows,{'raw_rows':len(rows),'duplicate_rows_removed':duplicate,'admissions':len(allrows),'people':len(first),'excluded_repeat_admissions':len(allrows)-len(first),'structure':structure,'alerts':alerts}

def analyze(selected,audit):
    r={'flow':{k:v for k,v in audit.items() if k!='structure'},'variables':{k:summary([x[k] for x in selected]) for k in ['age','crp_first','crp_last','stay_days']}}
    pair=[x for x in selected if x['crp_first'] is not None and x['crp_last'] is not None]
    delta=[x['crp_last']-x['crp_first'] for x in pair]
    count=Counter('decreased' if d<0 else 'increased' if d>0 else 'unchanged' for d in delta)
    r['paired']={**summary(delta),**{k:count[k] for k in ['decreased','increased','unchanged']}}
    ox=[x['oxygen'] for x in selected if x['oxygen'] is not None]; n=len(ox); k=sum(ox); z=1.959963984540054
    if n:
        p=k/n; mid=(p+z*z/(2*n))/(1+z*z/n); h=z*math.sqrt(p*(1-p)/n+z*z/(4*n*n))/(1+z*z/n)
        r['oxygen']={'n':n,'yes':k,'percent':100*p,'ci_low':100*(mid-h),'ci_high':100*(mid+h)}
    else: r['oxygen']={'n':0,'yes':0,'percent':None,'ci_low':None,'ci_high':None}
    assert sum(count.values())==len(pair)
    return r,pair

def legacy_plots(selected,pair,out):
    """One coordinate source for SVG and PNG, no dependency on plotting framework."""
    from PIL import Image,ImageDraw,ImageFont
    try: font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',18)
    except OSError: font=ImageFont.load_default()
    outputs=[]
    for name,points,xlab,ylab in [
        ('figure-1',[(float(x['crp_first']),float(x['stay_days']),x['person']) for x in selected if x['crp_first'] is not None and x['stay_days'] is not None],'First recorded CRP (mg/L)','Recorded stay (days)'),
        ('figure-2',[(float(x['crp_first']),float(x['crp_last']),x['person']) for x in pair],'First recorded CRP (mg/L)','Last recorded CRP (mg/L)')]:
        im=Image.new('RGB',(900,560),'white'); d=ImageDraw.Draw(im)
        svg=['<svg xmlns="http://www.w3.org/2000/svg" width="900" height="560" viewBox="0 0 900 560"><rect width="900" height="560" fill="white"/>']
        def line(coords,color='black'):
            d.line(coords,fill=color,width=2); svg.append(f'<line x1="{coords[0]}" y1="{coords[1]}" x2="{coords[2]}" y2="{coords[3]}" stroke="{color}" stroke-width="2"/>')
        def text(x,y,s):
            d.text((x,y),s,fill='#243247',font=font); svg.append(f'<text x="{x}" y="{y+17}" font-family="Arial,sans-serif" font-size="18" fill="#243247">{s}</text>')
        mx=max([p[0] for p in points]+[1])*1.12; my=max([p[1] for p in points]+[1])*1.12
        if name=='figure-2': mx=my=max(mx,my)
        line((100,440,820,440)); line((100,440,100,70))
        for i in range(6):
            xx=100+i*144; yy=440-i*74
            line((xx,440,xx,447)); text(xx-15,452,fmt(mx*i/5)); line((94,yy,100,yy)); text(35,yy-12,fmt(my*i/5))
        text(250,510,xlab); text(100,15,f'{ylab} | first admission per person | n={len(points)}')
        if name=='figure-2': line((100,440,820,70),'#b8c3d1')
        for x,y,_ in points:
            px=100+x/mx*720; py=440-y/my*370
            d.ellipse((px-5,py-5,px+5,py+5),fill='#247b98',outline='white'); svg.append(f'<circle cx="{px:.6f}" cy="{py:.6f}" r="5" fill="#247b98" stroke="white"/>')
        text(100,480,'Overlapping values may share a point; no treatment-effect interpretation.')
        (out/(name+'.svg')).write_text('\n'.join(svg)+ '</svg>',encoding='utf8'); im.save(out/(name+'.png'))
        save(out/(name+'-data.json'),{'points':points,'x_unit':'mg/L','y_unit':'days' if name=='figure-1' else 'mg/L','x_max':mx,'y_max':my})
        outputs.append(name)
    return outputs

def display_values(r):
    f=r['flow']; p=r['paired']; o=r['oxygen']; v=r['variables']; n=f['people']
    age=v['age']; crp=v['crp_first']; days=v['stay_days']
    values={'n':str(n),'rows':str(f['raw_rows']),'dup':str(f['duplicate_rows_removed']),'admissions':str(f['admissions']),'repeat':str(f['excluded_repeat_admissions']),
        'age_n':str(age['n']),'age_mean':fmt(age['mean']),'age_median':fmt(age['median']),
        'crp_n':str(crp['n']),'crp_median':fmt(crp['median']),'crp_q1':fmt(crp['q1']),'crp_q3':fmt(crp['q3']),
        'days_n':str(days['n']),'days_median':fmt(days['median']),'days_q1':fmt(days['q1']),'days_q3':fmt(days['q3']),
        'pair_n':str(p['n']),'delta_median':fmt(p['median']),'delta_q1':fmt(p['q1']),'delta_q3':fmt(p['q3']),
        'decreased':str(p['decreased']),'increased':str(p['increased']),'unchanged':str(p['unchanged']),
        'ox_n':str(o['n']),'ox_yes':str(o['yes']),'ox_pct':fmt(o['percent']),'ox_lo':fmt(o['ci_low']),'ox_hi':fmt(o['ci_high'])}
    return values

def documents(r,c,refs,out):
    """Constrained template authoring: every displayed research number from this result."""
    values=display_values(r); v=r['variables']
    # Each template placeholder remains an inspectable structured claim.
    save(out/'display-values.json',values)
    text={
    'zh':[
    ('摘要', '目的：描述合成住院资料的变量分布及复测子集变化，验证从脏表到可审阅论文的流程。方法：以每人首次住院为单位，去除完全重复行、统一检验单位，按可用病例描述，不插补。结果：原表{{rows}}行，去重{{dup}}行后有{{admissions}}次住院；纳入{{n}}人首次住院。首次记录CRP中位数{{crp_median}} mg/L（四分位数{{crp_q1}}至{{crp_q3}}；有效{{crp_n}}人）。复测{{pair_n}}人，下降{{decreased}}人、上升{{increased}}人、不变{{unchanged}}人，末减首中位数{{delta_median}} mg/L。结论：本例能形成可追溯描述，复测变化不代表治疗效果；无随访，不能推断远期结局。本稿为合成数据演示，不是可投稿的真实研究。'),
    ('引言', '临床表格常同时包含重复住院、缺失检验及不一致单位。若不先说明分析单位，记录数会被误当人数。本演示研究聚焦“现有资料能描述什么”，不使用显著性筛选结果。观察性研究的透明报告可参考STROBE，但报告规范本身不证明研究质量[1]。'),
    ('引言', '本例的目标是'+c['objective']['zh']+'。后续分析围绕这一描述目标，不据图中分布推断预后价值。'),
    ('方法', '设计与来源：合成住院资料的回顾性描述演示。样本规模由演示生成器设定，不是基于临床效应或检验效能确定。没有真实招募、伦理批准、机构或作者信息。'+c['population']['zh']),
    ('方法', '数据整理：按明确的人和住院编号处理，完全相同行只保留一行；冲突编号不自动合并。按入院日期选每人首次住院，同日冲突须澄清。CRP统一为mg/L，mg/dL乘以10；用十进制计算升降，避免等值换算后的浮点尾差。首末值是表内首次与末次记录，不预设为治疗前后。'),
    ('方法', '统计：连续变量报告有效人数、中位数及线性插值四分位数，年龄另报告均值；氧疗为已知编码中的人数/分母，区间使用Wilson方法。CRP配对限定首末均有记录者，变化为末次减首次；缺失不当作零变化。未做插补、假设检验、调整模型或多重比较。散点用于观察记录分布，不作因果或预测效能判断。'),
    ('方法', '随访没有，也补不到，已在模拟确认卡中记录一次；不重复追要，不当作无事件。表载死亡后再入院矛盾仅列质控警示，死亡标签未核，本分析不涉及死亡或生存。真实常规医疗数据研究还应交代数据来源与选择过程[2]；本演示不宣称完整遵循其所有条目。'),
    ('结果', '流程与缺失：原表{{rows}}行，移除{{dup}}行完全重复后有{{admissions}}次住院，首次住院队列{{n}}人，额外住院{{repeat}}次不作为独立人纳入。年龄有效{{age_n}}人，均值{{age_mean}}岁、中位数{{age_median}}岁。住院记录天数有效{{days_n}}人，中位数{{days_median}}天（{{days_q1}}至{{days_q3}}）。这是表内日数，不是康复或生存时间。见表1。'),
    ('结果', '首次记录CRP有效{{crp_n}}人，中位数{{crp_median}} mg/L（{{crp_q1}}至{{crp_q3}}）。氧疗在{{ox_n}}名编码已知者中记录{{ox_yes}}人，比例{{ox_pct}}%，95% Wilson区间{{ox_lo}}%至{{ox_hi}}%；未编码者不计为无氧疗。图1仅展示CRP和日数都有记录的首次住院。'),
    ('结果', 'CRP完整配对{{pair_n}}人：下降{{decreased}}人、上升{{increased}}人、不变{{unchanged}}人；变化中位数{{delta_median}} mg/L（{{delta_q1}}至{{delta_q3}}）。方向三类之和等于配对人数。图2与表2采用同一配对集合，不能把缺复测者算成不变。'),
    ('讨论', '本演示的主要发现是资料可用范围及复测异质性：{{pair_n}}名复测者中下降{{decreased}}人，上升{{increased}}人，不变{{unchanged}}人。这些变化来自被选择复测的子集，既没有未复测者变化，也没有治疗暴露或可比较对照，不能归因于治疗或代表全队列改善。'),
    ('讨论', '本例把有效分母、排除规则及未知来源写明，采用报告规范的透明表达原则[1,2]。但这与真实疾病队列的临床证据不是同一层面；未检索疾病特异研究、未评估诊断价值，不能据本例推荐CRP阈值或临床路径。图上可见的分布也不能代替正式关联方案。'),
    ('局限', '合成数据没有真实临床外部效度。样本小、单位混杂和缺失规则由生成器人为安排；复测选择可能改变分布。首末记录时间与治疗时点未建立关系，日数只来自现有记录，死亡矛盾未核、远期随访不可得。缺少混杂因素、纳排的真实执行证据及外部验证；因此不能证明因果、等效、疗效或预测能力。'),
    ('结论', '这份合成示例贯通了表格整理、有限描述、图表和中英文稿件。首末CRP变化限于{{pair_n}}名配对者；任何真实研究须另核定义、伦理和研究方案。本稿仅供流程审阅，不能以真实临床研究投稿。'),
    ('声明', '伦理、资金、机构、作者及利益冲突：真实研究信息未提供，不予虚构。合成输入与生成脚本随演示保存；本地处理，没有数据外发。参考文献身份与支持主张分别记录。')],
    'en':[
    ('Abstract', 'Objective: To describe synthetic admission records and repeat-measurement patterns, and demonstrate an auditable workflow from a messy spreadsheet to a manuscript. Methods: The first admission per person was selected after removing exact duplicates and harmonizing units. Available cases were summarized without imputation. Results: Of {{rows}} input rows, {{dup}} exact duplicates were removed, leaving {{admissions}} admissions and {{n}} first admissions. First recorded CRP had a median of {{crp_median}} mg/L (quartiles {{crp_q1}} to {{crp_q3}}; n={{crp_n}}). Among {{pair_n}} paired records, {{decreased}} decreased, {{increased}} increased and {{unchanged}} were unchanged; median last-minus-first change was {{delta_median}} mg/L. Conclusions: The workflow supports traceable description. Repeat measurements do not establish treatment effects; unavailable follow-up precludes long-term outcome analysis. This is a synthetic demonstration, not a real study suitable for submission.'),
    ('Introduction', 'Clinical spreadsheets may contain repeat admissions, missing measurements and mixed units. Without an explicit analysis unit, record counts can be mistaken for participant counts. This demonstration asks what the available records can describe, without selecting results for statistical significance. STROBE offers reporting guidance but does not establish study quality [1].'),
    ('Introduction', 'The objective was '+c['objective']['en']+'. The analysis addresses this descriptive objective; visible distributions do not establish prognostic value.'),
    ('Methods', 'Design and source: A retrospective-style descriptive demonstration of synthetic admission records. The sample size was chosen by the demonstration generator, not a clinical effect-size or power calculation. There was no real recruitment, ethics approval, institution or authorship information. '+c['population']['en']),
    ('Methods', 'Data preparation: Records were handled using explicit person and admission identifiers. Exact duplicates were removed; conflicting identifiers require clarification. The earliest dated admission was selected for each person; same-day ambiguity blocks automatic selection. CRP was standardized to mg/L by multiplying mg/dL by ten. Decimal arithmetic preserved equivalence when counting directions of change. First and last values denote recorded values, not assumed pre-treatment and post-treatment measurements.'),
    ('Methods', 'Analysis: Continuous variables were summarized using valid counts, medians and linearly interpolated quartiles; mean age was also reported. Oxygen use was summarized among known codes with a Wilson interval. Paired CRP analysis required both measurements and used last minus first. Missing repeat values were not assigned zero change. No imputation, hypothesis tests, adjusted models or multiple comparisons were performed. Scatterplots depict recorded distributions and do not estimate causal effects or predictive performance.'),
    ('Methods', 'Follow-up was unavailable and could not be supplemented, as recorded once in the simulated confirmation card. Absence of follow-up was not coded as absence of events. A death-labelled record followed by another admission was retained as a quality alert; death labels were unverified and no mortality or survival analysis was performed. Real routinely collected health-data studies should describe their sources and selection processes [2]; this demonstration does not claim complete compliance with all RECORD items.'),
    ('Results', 'Flow and missingness: Of {{rows}} input rows, {{dup}} exact duplicates were removed, leaving {{admissions}} admissions. The first-admission cohort included {{n}} people; {{repeat}} additional admissions were not treated as independent participants. Age was available for {{age_n}} people, with mean {{age_mean}} years and median {{age_median}} years. Recorded stay was available for {{days_n}} people, with median {{days_median}} days (quartiles {{days_q1}} to {{days_q3}}). These recorded durations do not denote recovery or survival time. See Table 1.'),
    ('Results', 'First recorded CRP was available for {{crp_n}} people: median {{crp_median}} mg/L (quartiles {{crp_q1}} to {{crp_q3}}). Oxygen use was recorded for {{ox_yes}} of {{ox_n}} people with known codes ({{ox_pct}}%; 95% Wilson interval {{ox_lo}}% to {{ox_hi}}%). Unknown codes were not treated as no oxygen use. Figure 1 includes first admissions with both CRP and recorded stay.'),
    ('Results', 'Complete CRP pairs were available for {{pair_n}} people: {{decreased}} decreased, {{increased}} increased and {{unchanged}} were unchanged. Median change was {{delta_median}} mg/L (quartiles {{delta_q1}} to {{delta_q3}}). Direction counts sum to the paired denominator. Figure 2 and Table 2 use the same paired set; missing repeats are not unchanged values.'),
    ('Discussion', 'The main findings describe data availability and heterogeneous repeat measurements: of {{pair_n}} paired records, {{decreased}} decreased, {{increased}} increased and {{unchanged}} were unchanged. These observations concern a selected retested subset. Changes in people without repeats are unknown. Neither treatment exposure nor a comparable control group was available. Consequently, changes cannot be attributed to treatment or interpreted as improvement across the full cohort.'),
    ('Discussion', 'Explicit denominators, exclusions and unknown provenance follow principles of transparent reporting [1,2]. They do not supply clinical evidence from an actual disease cohort. Disease-specific literature was not searched and diagnostic performance was not evaluated; no CRP threshold or care pathway is recommended. Visible scatterplot patterns cannot substitute for a prespecified association analysis. Any extension would require an appropriate question, definitions and a separately confirmed analysis plan.'),
    ('Limitations', 'Synthetic data lack clinical external validity. Sample size, missingness and unit irregularities were imposed by the generator. Selection for repeat measurements may alter distributions. The relation of recorded measurements to treatment timing is unknown. Stay was derived only from the recorded field; death contradictions remained unresolved and long-term follow-up was unavailable. Confounders, evidence of actual eligibility implementation and external validation were absent. These data cannot establish causation, equivalence, efficacy or prediction.'),
    ('Conclusions', 'The synthetic example links spreadsheet preparation, bounded description, figures and bilingual manuscripts. First-to-last CRP change applies only to {{pair_n}} paired records. Real research requires separate confirmation of definitions, ethics and design. This document demonstrates a workflow and must not be submitted as a real clinical study.'),
    ('Declarations', 'Ethics, funding, institution, authors and conflicts: Real-study information was not supplied and was not invented. Synthetic inputs and their generator are retained locally. No data were transmitted. Reference identity and claim support were documented separately.')]
    }
    claims=[]; docs={}; failures=[]
    for lang,sections in text.items():
        paragraphs=[]
        for i,(heading,tpl) in enumerate(sections):
            if i==0 or heading!=sections[i-1][0]: paragraphs.append({'text':heading,'evidence_ref':'section-'+str(i)})
            body=tpl
            for k in re.findall(r'\{\{(\w+)\}\}',tpl):
                body=body.replace('{{'+k+'}}',values[k]); claims.append({'language':lang,'section':heading,'paragraph':len(paragraphs),'template':tpl,'key':k,'display':values[k]})
            paragraphs.append({'text':body,'evidence_ref':'results' if '{{' in tpl else 'scope'})
        paragraphs.append({'text':'参考文献' if lang=='zh' else 'References','evidence_ref':'refs'})
        for ref in refs['references']: paragraphs.append({'text':f"[{ref['id']}] {ref['citation']} {ref['url']}",'evidence_ref':'ref-'+str(ref['id'])})
        columns=['指标','有效n','中位数','Q1','Q3','单位'] if lang=='zh' else ['Measure','Valid n','Median','Q1','Q3','Unit']
        labels={'age':('年龄','Age','years'),'crp_first':('首次CRP','First CRP','mg/L'),'crp_last':('末次CRP','Last CRP','mg/L'),'stay_days':('住院日数','Recorded stay','days')}
        rows=[[labels[k][0 if lang=='zh' else 1],str(x['n']),fmt(x['median']),fmt(x['q1']),fmt(x['q3']),({'years':'岁','days':'天'}.get(labels[k][2],labels[k][2]) if lang=='zh' else labels[k][2])] for k,x in v.items()]
        tab1={'caption':('表1 首次住院可用病例描述；缺失不插补' if lang=='zh' else 'Table 1. Available-case description of first admissions; no imputation'),'columns':columns,'rows':rows,'evidence_ref':'variables'}
        tab2={'caption':('表2 CRP完整配对：末减首；仅复测子集' if lang=='zh' else 'Table 2. Complete CRP pairs: last minus first; retested subset only'),'columns':(['配对n','下降','上升','不变','变化中位数 mg/L'] if lang=='zh' else ['Paired n','Decreased','Increased','Unchanged','Median change mg/L']),'rows':[[values[k] for k in ['pair_n','decreased','increased','unchanged','delta_median']]],'evidence_ref':'paired'}
        figs=[]
        for i in [1,2]:
            cap= [('图1 首次记录CRP与住院日数；完整记录；仅描述分布','Figure 1. First recorded CRP and stay; complete records; descriptive distribution only'),('图2 首末CRP完整配对；灰线为等值；不能解释为治疗效果','Figure 2. Complete first/last CRP pairs; grey line indicates equality; no treatment-effect interpretation')][i-1][0 if lang=='zh' else 1]
            cap+=('；相同坐标按面积表示重复次数，未移动数据' if lang=='zh' else '; identical coordinates encoded by marker area, without jitter')
            asset_path=f'graphics/{lang}/figure-{i}.png'
            figs.append({'caption':cap,'relative_path':asset_path,'sha256':sha(out/asset_path),'evidence_ref':'figure-'+str(i)})
        doc={'schema_version':'aggregate-materialization/1.0.0','source_package_ref':sha(out/'results.json'),'title':c['title'][lang],'paragraphs':paragraphs,'tables':[tab1,tab2],'figures':figs}
        save(out/f'document-{lang}.json',doc); docs[lang]=doc
        try: materialize.materialize(doc,out,out/lang)
        except (ImportError,OSError,subprocess.SubprocessError) as e: failures.append(e)
    save(out/'numeric-claims.json',claims)
    if failures: raise failures[0]
    return docs

def reference_check(refs):
    allowed={
        'C1':('10.1371/journal.pmed.0040296','STROBE指导透明报告，不证明研究质量','STROBE offers reporting guidance, not proof of study quality'),
        'C2':('10.1371/journal.pmed.1001885','真实常规医疗数据研究应报告来源与选择过程','Real routinely collected health-data studies should describe sources and selection')}
    refmap={r['id']:r for r in refs['references']}
    if set(x['id'] for x in refs['claims'])!=set(allowed): raise Blocked('当前固定方法稿只接受已核C1/C2；其他主张需另做内容核证。')
    for claim in refs['claims']:
        doi,zh,en=allowed[claim['id']]
        if refmap[claim['reference_id']].get('doi')!=doi or claim.get('zh')!=zh or claim.get('en')!=en:
            raise Blocked('真实文献不能支持新临床主张；当前内容支持合同不匹配，需要逐项核证。')
    ids=[r['id'] for r in refs['references']]
    if ids!=[1,2]: raise Blocked('当前自动描述稿需经核的STROBE/RECORD方法文献，不能静默替换。')
    for r in refs['references']:
        if not r.get('metadata_evidence') or not r.get('content_evidence') or not r.get('verified_date'): raise Blocked('引用缺少元数据或内容支持证据。')
    for x in refs['claims']:
        if x['status']!='supported' or not x.get('evidence_location'): raise Blocked('文献存在但主张未支持；先改写/删除或核证。')

def cell_diff(previous,current):
    old={x['admission']:x for x in previous}; new={x['admission']:x for x in current}; changed=[]
    for a in sorted(old.keys()|new.keys()):
        for k in sorted(set(old.get(a,{}))|set(new.get(a,{}))):
            before=old.get(a,{}).get(k); after=new.get(a,{}).get(k)
            if before!=after: changed.append({'admission':a,'field':k,'before':before,'after':after})
    return changed

def serial_rows(rows): return [{k:str(v) if isinstance(v,Decimal) else v for k,v in x.items()} for x in rows]

def run(path,c,refs_path,out,previous=None,analysis_only=False,gate=None,case=None,evidence_root=None,question=None,activity="DESCRIPTIVE"):
    permission=require_gate(gate,case,evidence_root,path,question,activity if analysis_only else "DRAFT",synthetic_demo=c.get("synthetic") is True)
    if (permission["case_kind"]=="SYNTHETIC") != (c.get("synthetic") is True): raise Blocked("Gate/config synthetic identity mismatch")
    if activity=="CONFIRMATORY": raise Blocked("Fixed engine implements descriptive analysis only; use a confirmed host method for effects")
    out=Path(out)
    if out.exists(): raise Blocked('新输出目录已存在；不覆写旧版。失败后请核记录，换新目录重试。')
    if previous:
        oldc=load(Path(previous)/'config.json')
        if scientific_contract(oldc)!=scientific_contract(c): raise Blocked('研究定义/方法发生变化：需新的医生科学确认，不能当作单纯改表更新。')
    if not analysis_only:
        if c.get('synthetic') is not True: raise Blocked('固定自动全文是合成演示模板；真实数据请用 analyze 后由宿主按SKILL撰写，经通用物化接口交付，不套用演示事实。')
        if c.get('followup',{}).get('status')!='unavailable' or c.get('followup',{}).get('can_supplement') is not False:
            raise Blocked('固定演示稿要求已有“随访没有且补不到”的确认；其他情况用 analyze 和宿主撰稿。')
        refs=load(refs_path); reference_check(refs)
    else: refs={'scope':'No manuscript references needed for analysis-only mode.'}
    selected,allrows,audit=clean(path,c)
    out.mkdir(parents=True); save(out/'config.json',c); save(out/'references.json',refs)
    # Preserve exact source bytes and provenance locally, separate from aggregate manuscript.
    source=out/('input'+Path(path).suffix.lower()); source.write_bytes(Path(path).read_bytes())
    save(out/'input-manifest.json',{'source_filename':Path(path).name,'source_sha256':sha(path),'snapshot':source.name,'synthetic':c.get('synthetic',False)})
    save(out/'cleaned-records.local.json',serial_rows(allrows)); save(out/'selected-records.local.json',serial_rows(selected)); save(out/'quality.json',audit)
    r,pair=analyze(selected,audit); save(out/'results.json',r)
    save(out/'checkpoint.json',{'stage':'RESULTS_READY','source_sha256':sha(source),'config_sha256':sha(out/'config.json'),'results_sha256':sha(out/'results.json'),'selected_sha256':sha(out/'selected-records.local.json'),'cleaned_sha256':sha(out/'cleaned-records.local.json'),'quality_sha256':sha(out/'quality.json'),'next':'graphics; manuscript not generated'})
    v02plots(selected,pair,out,synthetic=c.get("synthetic",False))
    if analysis_only:
        (out/'医生首页.md').write_text(f'# 有限描述分析\n\n共{r["flow"]["people"]}人首次住院，CRP完整配对{r["paired"]["n"]}人。已生成实际图及同源结果。缺失保留，死亡矛盾不用于结局分析。\n\n本次没有自动全文；请按SKILL由宿主基于结果、已核文献和真实医生事实撰写，不使用合成演示事实。\n',encoding='utf8')
        save(out/'validation.json',{'status':'ANALYSIS_ONLY','manuscript':'NOT_GENERATED'}); return r
    documents(r,c,refs,out)
    if previous:
        prev=Path(previous); oldvalues=load(prev/'display-values.json'); newvalues=load(out/'display-values.json')
        changed=[{'key':k,'before':oldvalues[k],'after':v} for k,v in newvalues.items() if oldvalues[k]!=v]
        oldclaims=load(prev/'numeric-claims.json'); impacted=[x for x in oldclaims if x['key'] in {v['key'] for v in changed}]
        artifacts=[]
        for p in out.rglob('*'):
            if p.is_file() and (prev/p.relative_to(out)).is_file() and sha(p)!=sha(prev/p.relative_to(out)): artifacts.append(str(p.relative_to(out)))
        save(out/'correction.json',{'previous':str(prev.resolve()),'previous_source_hash':load(prev/'input-manifest.json')['source_sha256'],'current_source_hash':sha(source),'normalized_cell_changes':cell_diff(load(prev/'cleaned-records.local.json'),serial_rows(allrows)),'changed_display_values':changed,'impacted_prior_numeric_claims':impacted,'changed_artifacts':artifacts,'policy':'All current results, plots, bilingual sections and DOCX rebuilt; references unchanged unless evidence changes.'})
        (out/'更正说明.md').write_text('# 本次更正\n\n旧版完整保留；定义和方法未改变。新版重新清理、重算，重建两种语言与DOCX。\n\n'+ '\n'.join(f"- {x['key']}: {x['before']} → {x['after']}" for x in changed)+'\n\n完整受影响段落与文件见 correction.json。PDF在显式渲染后重新生成。\n',encoding='utf8')
    (out/'医生首页.md').write_text(f"# 这次能得到什么\n\n{'合成数据演示；不能以真实研究投稿。' if c.get('synthetic') else '真实数据描述草稿；尚需医生科学审阅。'}\n\n按每人首次住院，共{r['flow']['people']}人；CRP有{r['paired']['n']}人可配对，下降{r['paired']['decreased']}、上升{r['paired']['increased']}、不变{r['paired']['unchanged']}。无随访且补不到，本轮不再追问。\n\n现有资料支持有限描述及两张图，不能支持治疗效果、死亡或远期预后。表载死亡后入院矛盾{len(audit['alerts'])}项保留待核。\n\n中文稿：[DOCX](zh/manuscript.docx) · [Markdown](zh/manuscript.md)\n英文稿：[DOCX](en/manuscript.docx) · [Markdown](en/manuscript.md)\n\n所有结果来自同一results.json。定义改变需再次科学确认；纯数据更正可沿本方案重跑并保留历史。\n",encoding='utf8')
    save(out/'validation.json',verify(out)); return r

def verify(out):
    out=Path(out); c=load(out/'config.json'); source=load(out/'input-manifest.json'); assert sha(out/source['snapshot'])==source['source_sha256']
    selected,allrows,audit=clean(out/source['snapshot'],c); r,pair=analyze(selected,audit)
    assert r==load(out/'results.json'); claims=load(out/'numeric-claims.json'); values=load(out/'display-values.json')
    assert values==display_values(r), 'Display values detached from result source'
    for lang in ['zh','en']:
        doc=load(out/f'document-{lang}.json'); materialize.qc(doc,out/lang/'manuscript.docx')
        for claim in [x for x in claims if x['language']==lang]:
            assert claim['display']==values[claim['key']]
            expected=claim['template']
            for key in re.findall(r'\{\{(\w+)\}\}',expected): expected=expected.replace('{{'+key+'}}',values[key])
            assert doc['paragraphs'][claim['paragraph']]['text']==expected, 'Numeric paragraph detached from results'
        assert not re.search(r'\{\{\w+\}\}',(out/lang/'manuscript.md').read_text())
    f1=load(out/'figure-1-data.json')['points']; f2=load(out/'figure-2-data.json')['points']
    assert f1==[ [float(x['crp_first']),float(x['stay_days']),x['person']] for x in selected if x['crp_first'] is not None and x['stay_days'] is not None]
    assert f2==[[float(x['crp_first']),float(x['crp_last']),x['person']] for x in pair]
    for name,data in [('figure-1',f1),('figure-2',f2)]:
        for lang in ['en','zh']:
            folder=out/'graphics'/lang; payload=load(folder/(name+'.json')); plotted=payload['data']
            assert plotted['x']==[p[0] for p in data] and plotted['y']==[p[1] for p in data]
            counts=Counter((p[0],p[1]) for p in data)
            assert plotted['multiplicity']==[counts[k] for k in sorted(counts)] and sum(plotted['multiplicity'])==len(data)
            assert all(payload['audit'][k]=='pass' for k in ['glyphs','canvas_bounds','tick_overlap'])
            for filename,digest in payload['files'].items(): assert sha(folder/filename)==digest
        assert sha(out/(name+'.png'))==sha(out/'graphics/en'/(name+'.png'))
    reference_check(load(out/'references.json'))
    return {'status':'PASS_BOUNDED','recomputed_results':True,'docx_exact_text_tables_figures':True,'numeric_claim_occurrences':len(claims),'plot_points':{'figure-1':len(f1),'figure-2':len(f2)},'reference_contract':'PASS (offline evidence records; not independent semantic review)','visual':'REQUIRES_RENDER_AND_INSPECTION','scientific_review':'NOT_PERFORMED','generic_manuscript_authoring':'NOT_IMPLEMENTED; constrained descriptive templates only'}

def main():
    if not __debug__: raise SystemExit('Verification requires normal Python; do not use -O.')
    a=argparse.ArgumentParser(description=__doc__); sub=a.add_subparsers(dest='cmd',required=True)
    for op in ['inspect','run','analyze']:
        p=sub.add_parser(op); p.add_argument('input',type=Path); p.add_argument('--config',type=Path,required=True); p.add_argument('--out',type=Path,required=True)
        if op in {'inspect','run','analyze'}:
            for flag in ['gate','case','evidence-root','question']: p.add_argument('--'+flag,required=True)
            p.add_argument('--activity',choices=['DESCRIPTIVE','CONFIRMATORY'],default='DESCRIPTIVE')
        if op in {'run','analyze'}: p.add_argument('--references',type=Path,required=op=='run'); p.add_argument('--previous',type=Path)
    p=sub.add_parser('verify'); p.add_argument('output',type=Path)
    p=sub.add_parser('doctor'); p.add_argument('--request',choices=[*environment.CAPS,'pdf','all'],default='all'); p.add_argument('--runtime-root')
    args=a.parse_args()
    try:
        if args.cmd=='doctor':
            report=environment.check(args.request,args.runtime_root)
            print(json.dumps(report,ensure_ascii=False,indent=2)); return 0 if report['status']=='READY' else 3
        if args.cmd=='verify': print(json.dumps(verify(args.output),ensure_ascii=False,indent=2)); return 0
        c=load(args.config)
        if args.cmd=='inspect':
            require_gate(args.gate,args.case,args.evidence_root,args.input,args.question,'QC')
            inspect(args.input,c,args.out)
        else: run(args.input,c,args.references,args.out,args.previous,args.cmd=='analyze',args.gate,args.case,args.evidence_root,args.question,args.activity)
        print('完成：'+str(args.out)); return 0
    except (ImportError,OSError,subprocess.SubprocessError) as e:
        return environment.failure(e,getattr(args,'out',None))
    except (Blocked,AssertionError,KeyError,ValueError) as e:
        print('暂停：'+str(e),file=sys.stderr); return 2
if __name__=='__main__': raise SystemExit(main())

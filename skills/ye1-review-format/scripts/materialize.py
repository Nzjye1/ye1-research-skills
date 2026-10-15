"""Deterministic aggregate manuscript materialization; never computes research results."""
from __future__ import annotations

# All data-bearing streams stay behind the controlled local adapter.
if __name__ == '__main__':
    import sys as _privacy_sys
    if '--help' not in _privacy_sys.argv[1:]:
        from privacy_boundary import cli_local_entry
        cli_local_entry()

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile

class MaterializationBlocked(ValueError): pass
class ToolUnavailable(MaterializationBlocked, OSError): pass
TOKEN=re.compile(r"[A-Za-z0-9_.-]+")
NUMBER=re.compile(r"(?<!\w)[+-]?\d+(?:\.\d+)?(?:%|e[+-]?\d+)?",re.I)

def evidence(value):
    if not isinstance(value,str) or not TOKEN.fullmatch(value):
        raise MaterializationBlocked("LOGICAL_EVIDENCE_REQUIRED")

def asset(root,relative,expected):
    if not isinstance(relative,str) or ":" in relative or "\\" in relative or any(p in {"",".",".."} for p in relative.split("/")):
        raise MaterializationBlocked("UNSAFE_ASSET_PATH")
    target=Path(root)
    for part in relative.split("/"):
        target/=part
        if target.is_symlink(): raise MaterializationBlocked("SYMLINK_ASSET")
    if not target.resolve().is_relative_to(Path(root).resolve()) or not target.is_file():
        raise MaterializationBlocked("MISSING_ASSET")
    if hashlib.sha256(target.read_bytes()).hexdigest()!=expected:
        raise MaterializationBlocked("ASSET_HASH_MISMATCH")
    return target

def validate(document):
    if set(document)-{"language","review_warnings"}!={"schema_version","source_package_ref","title","paragraphs","tables","figures"} or document["schema_version"]!="aggregate-materialization/1.0.0":
        raise MaterializationBlocked("DOCUMENT_SCHEMA")
    evidence(document["source_package_ref"])
    if not isinstance(document["title"],str) or not document["title"].strip(): raise MaterializationBlocked("TITLE_REQUIRED")
    for row in document["paragraphs"]:
        if set(row)!={"text","evidence_ref"} or not isinstance(row["text"],str): raise MaterializationBlocked("PARAGRAPH_SCHEMA")
        evidence(row["evidence_ref"])
    for table in document["tables"]:
        if set(table)!={"caption","columns","rows","evidence_ref"} or not table["columns"] or not isinstance(table["caption"],str): raise MaterializationBlocked("TABLE_SCHEMA")
        evidence(table["evidence_ref"])
        if any(not isinstance(cell,str) for cell in table["columns"]) or any(len(row)!=len(table["columns"]) or any(not isinstance(x,str) for x in row) for row in table["rows"]):
            raise MaterializationBlocked("TABLE_CELL_SCHEMA")
    for figure in document["figures"]:
        required={"caption","relative_path","sha256","evidence_ref"}
        if not required.issubset(figure) or set(figure)-required-{"display_width_cm"} or not isinstance(figure["caption"],str): raise MaterializationBlocked("FIGURE_SCHEMA")
        width=figure.get("display_width_cm",13)
        if not isinstance(width,(int,float)) or not 8<=width<=15: raise MaterializationBlocked("FIGURE_DISPLAY_WIDTH")
        evidence(figure["evidence_ref"])
    paragraphs=[x["text"] for x in document["paragraphs"]]
    for warning in document.get("review_warnings",[]):
        if set(warning)!={"anchor_text","text","evidence_ref"} or not all(isinstance(warning[k],str) and warning[k].strip() for k in ("anchor_text","text")):
            raise MaterializationBlocked("REVIEW_WARNING_SCHEMA")
        if warning["anchor_text"] not in paragraphs: raise MaterializationBlocked("REVIEW_WARNING_ANCHOR_MISSING")
        evidence(warning["evidence_ref"])
    if document.get("language") not in {None,"zh","en"}: raise MaterializationBlocked("LANGUAGE_ROUTE")
    language=document.get('language')
    labels=[x['caption'] for x in document['figures']+document['tables']]+[v for x in document['tables'] for v in x['columns']+[v for row in x['rows'] for v in row]]
    if language=='en' and any(re.search(r'[\u4e00-\u9fff]',v) for v in labels):raise MaterializationBlocked('EN_ASSET_LABEL_CONTAINS_CHINESE')
    if language=='zh' and any(re.search(r'\b(?:Table|Figure|Synthetic|Variable|Unit|Observed|Missing|Median|Complete|Permutation|days|years)\b',v,re.I) for v in labels):raise MaterializationBlocked('ZH_ASSET_LABEL_UNTRANSLATED')
    return document

def expected_text(document):
    items=[document["title"],*[p["text"] for p in document["paragraphs"]]]
    for table in document["tables"]:
        items += [table["caption"],*table["columns"],*[x for row in table["rows"] for x in row]]
    items += [figure["caption"] for figure in document["figures"]]
    items += [warning["text"] for warning in document.get("review_warnings",[])]
    return items

HEADINGS={'摘要','引言','方法','结果','讨论','局限','结论','声明','参考文献','Abstract','Introduction','Methods','Results','Discussion','Limitations','Conclusions','Declarations','References'}

def placement(document):
    """Insert each asset after its first body reference. Missing references block review."""
    plan=[]; placed=set()
    warnings_by_anchor={}
    for warning in document.get('review_warnings',[]):warnings_by_anchor.setdefault(warning['anchor_text'],[]).append(warning)
    for p in document['paragraphs']:
        plan.append(('paragraph',p))
        for warning in warnings_by_anchor.get(p['text'],[]):plan.append(('warning',warning))
        if p['text'] in HEADINGS: continue
        for kind, label in [('tables',r'(?:表|Table\s*)'),('figures',r'(?:图|Figure\s*)')]:
            for i,item in enumerate(document[kind],1):
                if (kind,i) not in placed and re.search(label+str(i)+r'(?!\d)',p['text']):
                    plan.append((kind,item));placed.add((kind,i))
    if len(placed)!=len(document['tables'])+len(document['figures']):
        raise MaterializationBlocked('MISSING_BODY_ASSET_REFERENCE')
    return plan

def qc(document,docx_path):
    from docx import Document
    actual=Document(docx_path)
    paragraphs=[x.text for x in actual.paragraphs if x.text]
    expected=[document['title']]+[(('AUTHOR REVIEW REQUIRED: '+x['text']) if kind=='warning' else x['text']) if kind in {'paragraph','warning'} else x['caption'] for kind,x in placement(document)]
    if paragraphs!=expected: raise MaterializationBlocked('DOCX_TEXT_MISMATCH')
    sources=[x for kind,x in placement(document) if kind=='tables']
    if len(actual.tables)!=len(sources): raise MaterializationBlocked('DOCX_TABLE_COUNT')
    cells=[]
    for table,source in zip(actual.tables,sources):
        rows=[[cell.text for cell in row.cells] for row in table.rows]
        if rows!=[source['columns'],*source['rows']]: raise MaterializationBlocked('DOCX_TABLE_MISMATCH')
        cells += [x for row in rows for x in row]
    if len(actual.inline_shapes)!=len(document['figures']): raise MaterializationBlocked('DOCX_FIGURE_COUNT')
    if Counter(NUMBER.findall('\n'.join(paragraphs+cells)))!=Counter(NUMBER.findall('\n'.join(expected_text(document)))):
        raise MaterializationBlocked('NUMERIC_TOKEN_MISMATCH')
    return {'status':'MATCH','verification_scope':'MATERIALIZATION_TECHNICAL','text_and_cells_exact':True,
            'figures_embedded':True,'numeric_tokens_match':True,'inline_at_first_reference':True,
            'visual_qc':'NOT_PERFORMED','scientific_acceptance':False}

def materialize_text(document,assets_root,output_root):
    document=validate(document); plan=placement(document)
    figures={id(x):asset(assets_root,x['relative_path'],x['sha256']) for x in document['figures']}
    output=Path(output_root)
    if output.exists(): raise MaterializationBlocked('OUTPUT_ALREADY_EXISTS')
    output.mkdir(parents=True)
    md=['# '+document['title']];idx=0
    for kind,item in plan:
        if kind=='paragraph':md.append(('## ' if item['text'] in HEADINGS else '')+item['text'])
        elif kind=='tables':
            md += [item['caption'],'| '+' | '.join(item['columns'])+' |','| '+' | '.join(['---']*len(item['columns']))+' |']
            md += ['| '+' | '.join(row)+' |' for row in item['rows']]
        elif kind=='warning':md.append('> **AUTHOR REVIEW REQUIRED:** '+item['text'])
        else:
            idx+=1;path=figures[id(item)];target=output/'figures'/f'figure-{idx}{path.suffix}';target.parent.mkdir(exist_ok=True);target.write_bytes(path.read_bytes())
            md.append(f"![{item['caption']}](figures/{target.name})")
    (output/'manuscript.md').write_text('\n\n'.join(md)+'\n',encoding='utf8')
    checkpoint={'stage':'TEXT_READY','source_package_ref':document['source_package_ref'],'text_sha256':hashlib.sha256((output/'manuscript.md').read_bytes()).hexdigest(),'next':'DOCX; REVIEW_PENDING'}
    (output/'checkpoint.json').write_text(json.dumps(checkpoint,indent=2))
    return checkpoint

def materialize(document,assets_root,output_root):
    checkpoint=materialize_text(document,assets_root,output_root)
    plan=placement(document);output=Path(output_root)
    from font_support import select
    selected_font=select(' '.join(expected_text(document)))
    (output/'font-preflight.json').write_text(json.dumps(selected_font,ensure_ascii=False,indent=2))
    figures={id(x):asset(assets_root,x['relative_path'],x['sha256']) for x in document['figures']}
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    article=Document(); cn=document.get('language','zh' if re.search(r'[\u4e00-\u9fff]',document['title']) else 'en')=='zh'
    caption_east_asia=selected_font['family']
    caption_bold_compatible=True
    # The macOS Arial Unicode MS face available to this workflow is Regular-only.
    # Use the installed Heiti SC Medium face for Chinese captions so inherited
    # bold styling resolves to a real East Asian weight instead of a missing glyph.
    if cn and selected_font['family']=='Arial Unicode MS':
        heiti_medium=Path('/System/Library/Fonts/STHeiti Medium.ttc')
        if sys.platform=='darwin' and heiti_medium.is_file():
            caption_east_asia='Heiti SC'
        else:
            caption_bold_compatible=False
    caption_font_preflight={**selected_font,
                            'caption_east_asia_family':caption_east_asia,
                            'caption_bold_compatible':caption_bold_compatible}
    (output/'font-preflight.json').write_text(json.dumps(caption_font_preflight,ensure_ascii=False,indent=2))
    for border in list(article.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    def page_layout(section,landscape=False):
        section.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width=Cm(29.7 if landscape else 21);section.page_height=Cm(21 if landscape else 29.7)
        section.top_margin=section.bottom_margin=Cm(1.6 if landscape else 2.3)
        section.left_margin=section.right_margin=Cm(1.7 if landscape else 2.4)
    page_layout(article.sections[0])
    for name,size in [('Normal',11),('Title',18),('Heading 1',14),('Caption',10)]:
        style=article.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+key),None)
        style_east_asia=caption_east_asia if name=='Caption' else selected_font['family']
        for key,font in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',style_east_asia),('cs','Arial')]:fonts.set(qn('w:'+key),font)
        pf=style.paragraph_format;pf.widow_control=True;pf.line_spacing=1.35;pf.space_after=Pt(6)
        if name in {'Title','Heading 1','Caption'}:pf.keep_with_next=True
    article.styles['Normal'].paragraph_format.first_line_indent=Cm(.74 if cn else 0)
    article.styles['Title'].paragraph_format.space_after=Pt(14)
    article.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    article.styles['Caption'].paragraph_format.first_line_indent=Cm(0)
    if not caption_bold_compatible:article.styles['Caption'].font.bold=False
    article.core_properties.author='';article.core_properties.last_modified_by=''
    article.add_heading(document['title'],0);md=['# '+document['title']];output.mkdir(parents=True,exist_ok=True)
    fig_index=0;abstract=False;refs=False
    for kind,item in plan:
        if kind=='paragraph':
            text=item['text']
            if text in HEADINGS:
                article.add_heading(text,1);md.append('## '+text);abstract=text in {'摘要','Abstract'};refs=text in {'参考文献','References'}
            else:
                p=article.add_paragraph(text);md.append(text)
                if abstract or refs:p.paragraph_format.first_line_indent=Cm(0)
                if refs:p.paragraph_format.space_after=Pt(4)
        elif kind=='tables':
            wide=len(item['columns'])>=6
            if wide:page_layout(article.add_section(WD_SECTION.NEW_PAGE),True)
            article.add_paragraph(item['caption'],'Caption');md.append(item['caption'])
            table=article.add_table(rows=1,cols=len(item['columns']));table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=True
            usable_width=25.0 if wide else 16.2
            for column in table.columns:column.width=Cm(usable_width/len(item['columns']))
            hdr=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(hdr)
            for c,text in zip(table.rows[0].cells,item['columns']):c.text=text
            for row in item['rows']:
                for c,text in zip(table.add_row().cells,row):c.text=text
            for row_i,row in enumerate(table.rows):
                tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
                for c in row.cells:
                    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for p in c.paragraphs:
                        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.15
                        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.keep_with_next=len(table.rows)<=8 and row_i<len(table.rows)-1
                        for run in p.runs:run.font.size=Pt(7.4 if wide else 8.8)
            md+=['| '+' | '.join(item['columns'])+' |','| '+' | '.join(['---']*len(item['columns']))+' |']
            md+=['| '+' | '.join(row)+' |' for row in item['rows']]
            if wide:page_layout(article.add_section(WD_SECTION.NEW_PAGE),False)
        elif kind=='warning':
            p=article.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.left_indent=Cm(.25);p.paragraph_format.right_indent=Cm(.25);p.paragraph_format.keep_together=True;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(9)
            ppr=p._p.get_or_add_pPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'FFF2CC');ppr.append(shd)
            borders=OxmlElement('w:pBdr');left=OxmlElement('w:left');left.set(qn('w:val'),'single');left.set(qn('w:sz'),'18');left.set(qn('w:space'),'6');left.set(qn('w:color'),'C65911');borders.append(left);ppr.append(borders)
            run=p.add_run('AUTHOR REVIEW REQUIRED: '+item['text']);run.bold=True;md.append('> **AUTHOR REVIEW REQUIRED:** '+item['text'])
        else:
            fig_index+=1;path=figures[id(item)]
            p=article.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1
            p.add_run().add_picture(str(path),width=Cm(item.get('display_width_cm',13)))
            caption=article.add_paragraph(item['caption'],'Caption');caption.paragraph_format.keep_with_next=False;caption.paragraph_format.keep_together=True
            target=output/'figures'/f'figure-{fig_index}{path.suffix}';target.parent.mkdir(exist_ok=True);target.write_bytes(path.read_bytes())
            md.append(f"![{item['caption']}](figures/{target.name})")
    # Apply explicit run-level font slots to every body, table, and linked footer run.
    # This avoids Word/LibreOffice choosing a missing East Asian theme font for text runs.
    latin='Times New Roman';east_asia=selected_font['family']
    def set_run_fonts(run,east_asia_family=east_asia):
        rpr=run._element.get_or_add_rPr();rfonts=rpr.rFonts
        if rfonts is None:rfonts=OxmlElement('w:rFonts');rpr.insert(0,rfonts)
        for theme in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):rfonts.attrib.pop(qn('w:'+theme),None)
        for slot,font in [('ascii',latin),('hAnsi',latin),('eastAsia',east_asia_family),('cs','Arial')]:rfonts.set(qn('w:'+slot),font)
        lang=rpr.find(qn('w:lang'))
        if lang is None:lang=OxmlElement('w:lang');rpr.append(lang)
        lang.set(qn('w:val'),'en-US');lang.set(qn('w:eastAsia'),'zh-CN' if cn else 'en-US')
    for paragraph in article.paragraphs:
        paragraph_east_asia=caption_east_asia if paragraph.style.name=='Caption' else east_asia
        for run in paragraph.runs:set_run_fonts(run,paragraph_east_asia)
    for table in article.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:set_run_fonts(run)
    for section in article.sections:
        for part in (section.header,section.footer):
            for paragraph in part.paragraphs:
                for run in paragraph.runs:set_run_fonts(run)
    docx_path=output/'manuscript.docx';article.save(docx_path)
    (output/'manuscript.md').write_text('\n\n'.join(md)+'\n',encoding='utf-8')
    receipt=qc(document,docx_path);receipt['source_package_ref']=document['source_package_ref'];receipt['stage']='CUSTOMER_REVIEW_PENDING'
    receipt['font_preflight']=caption_font_preflight
    receipt['outputs']={x.name:hashlib.sha256(x.read_bytes()).hexdigest() for x in (docx_path,output/'manuscript.md')}
    checkpoint.update(stage='DOCX_READY',docx_sha256=hashlib.sha256(docx_path.read_bytes()).hexdigest(),next='PDF and visual review; CUSTOMER_REVIEW_PENDING')
    (output/'checkpoint.json').write_text(json.dumps(checkpoint,indent=2))
    (output/'materialization-qc.local.json').write_text(json.dumps(receipt,indent=2)+'\n')
    return receipt

def explicit_tool(mapping,key):
    import os
    entry=mapping.get(key,{})
    p=Path(entry.get("executable") or "")
    if not p.is_absolute() or not p.is_file() or not os.access(p,os.X_OK): raise ToolUnavailable("EXPLICIT_TOOL_REQUIRED: "+key)
    if not re.fullmatch(r"[a-f0-9]{64}",entry.get("sha256", "")) or hashlib.sha256(p.read_bytes()).hexdigest()!=entry["sha256"]:
        raise MaterializationBlocked("TOOL_IDENTITY_MISMATCH")
    return str(p)

def render(docx_path,mapping,output_root,*,runner=subprocess.run):
    if mapping.get('enabled') is not True:raise MaterializationBlocked('RENDER_DISABLED')
    explicit_tool(mapping,'soffice')
    from font_support import select,configured
    from zipfile import ZipFile, BadZipFile
    from xml.etree import ElementTree
    try:
        with ZipFile(docx_path) as archive:
            tree=ElementTree.fromstring(archive.read('word/document.xml'))
            text=' '.join(n.text or '' for n in tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
    except (BadZipFile,KeyError,ElementTree.ParseError) as error:raise MaterializationBlocked('Invalid DOCX source: '+str(error)) from error
    font=select(text)
    with configured(font):
        result=_render(docx_path,mapping,output_root,runner=runner)
    (Path(output_root)/'font-preflight.json').write_text(json.dumps(font,ensure_ascii=False,indent=2))
    return {**result,'font_preflight':font}

def _render(docx_path,mapping,output_root,*,runner=subprocess.run):
    if mapping.get("enabled") is not True: raise MaterializationBlocked("RENDER_DISABLED")
    soffice=explicit_tool(mapping,"soffice")
    output=Path(output_root)
    if output.exists(): raise MaterializationBlocked("RENDER_OUTPUT_EXISTS")
    output.mkdir(parents=True)
    with tempfile.TemporaryDirectory(prefix="docx-profile-",dir=output) as profile:
        runner([soffice,"-env:UserInstallation="+Path(profile).as_uri(),"--headless","--convert-to","pdf","--outdir",str(output),str(Path(docx_path).resolve())],check=True,timeout=120,capture_output=True)
    pdf=output/(Path(docx_path).stem+".pdf")
    if not pdf.is_file() or not pdf.read_bytes().startswith(b"%PDF-"): raise MaterializationBlocked("PDF_NOT_PRODUCED")
    (output/'checkpoint.json').write_text(json.dumps({'stage':'PDF_READY','pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'visual':'PENDING','next':'page images and visual review'}))
    poppler=explicit_tool(mapping,"pdftoppm")
    runner([poppler,"-png","-r","150",str(pdf),str(output/"page")],check=True,timeout=120,capture_output=True)
    pages=list(output.glob("page-*.png"))
    if not pages: raise MaterializationBlocked("PAGES_NOT_PRODUCED")
    from PIL import Image
    for p in pages:
        with Image.open(p) as img: img.verify()
    return {"status":"MATERIALIZED","verification_scope":"RENDER_TECHNICAL", "page_files_present":True,
            "visual_qc":"REQUIRES_INSPECTION","scientific_acceptance":False}

def main():
    import sys
    if len(sys.argv)>1 and sys.argv[1]=="render":
        parser=argparse.ArgumentParser(description="Explicit frozen DOCX render")
        parser.add_argument("docx",type=Path);parser.add_argument("--tools",required=True,type=Path)
        parser.add_argument("--output-root",required=True,type=Path)
        args=parser.parse_args(sys.argv[2:]);receipt=render(args.docx,json.loads(args.tools.read_bytes()),args.output_root)
        (args.output_root/"render-qc.local.json").write_text(json.dumps(receipt,indent=2)+"\n")
        return 0
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input",type=Path);parser.add_argument("--assets-root",required=True,type=Path)
    parser.add_argument("--output-root",required=True,type=Path);parser.add_argument("--text-only",action="store_true")
    args=parser.parse_args();fn=materialize_text if args.text_only else materialize;fn(json.loads(args.input.read_bytes()),args.assets_root,args.output_root)
    return 0
if __name__=="__main__":
    try: raise SystemExit(main())
    except (ImportError, OSError, subprocess.SubprocessError) as exc:
        print(json.dumps({'status':'BLOCKED_ENVIRONMENT','reason':str(exc),'next_step':'Retain Markdown/checkpoint already created; resume in an existing capable environment. Do not auto-install.'}),file=sys.stderr)
        raise SystemExit(3)
    except (ValueError, KeyError, TypeError) as exc:
        print(json.dumps({'status':'BLOCKED_VALIDATION','reason':str(exc)}),file=sys.stderr)
        raise SystemExit(2)
